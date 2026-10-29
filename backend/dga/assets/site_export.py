"""Asset-owned complete site reference workbook; no import or maintenance commands."""
from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal
from pathlib import Path
from tempfile import SpooledTemporaryFile
from typing import BinaryIO
from uuid import UUID, uuid4
from zipfile import ZipFile, ZIP_DEFLATED

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.writer.excel import ExcelWriter
from sqlalchemy import text

XLSX_TYPE = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
SCHEMA_VERSION = 'site-reference-v1'
COLUMNS = (
    ('site_id', '现场编号：既有UUID，文本'),
    ('site_name', '现场名称'),
    ('customer_id', '客户编号：既有UUID，文本'),
    ('customer_name', '客户名称'),
    ('location_text', '现场位置'),
    ('product_lines', '产品线：PV、ESS或PV;ESS；空白表示未配置'),
    ('pv_power_mw', '光伏装机功率，MW；十进制文本保留精度'),
    ('ess_power_mw', '储能装机功率，MW；十进制文本保留精度'),
    ('ess_energy_mwh', '储能电池容量，MWh；十进制文本保留精度'),
    ('commissioning_date', 'Commissioning Date，YYYY-MM-DD'),
    ('grid_status', '并网调试状态：NOT_STARTED=未开始；IN_PROGRESS=调试中；COMPLETED=已完成'),
    ('operation_status', '投运状态：NOT_OPERATIONAL=未投运；PARTIAL=部分投运；OPERATIONAL=已投运；DECOMMISSIONED=已退役'),
    ('source_note', '来源备注：当前未提供时保持空白，不推断'),
)
QUERY = text('''SELECT s.id AS site_id,s.site_name,s.customer_id,c.customer_name,s.location_text,s.source_note,
    pv.site_id IS NOT NULL AS has_pv,ess.site_id IS NOT NULL AS has_ess,
    pv.power_mw AS pv_power_mw,ess.power_mw AS ess_power_mw,ess.energy_mwh AS ess_energy_mwh,
    s.commissioning_date,s.grid_status,s.operation_status
    FROM sites s JOIN customers c ON c.id=s.customer_id
    LEFT JOIN site_product_lines pv ON pv.site_id=s.id AND pv.product_line='PV'
    LEFT JOIN site_product_lines ess ON ess.site_id=s.id AND ess.product_line='ESS'
    ORDER BY s.site_name,s.id''')


@dataclass(frozen=True)
class SiteExportFile:
    file: BinaryIO
    filename: str
    generated_at: datetime
    site_count: int
    export_id: UUID
    schema_version: str = SCHEMA_VERSION


class SiteExportUnavailable(Exception):
    pass


def _row(sheet, values):
    cells = []
    for value in values:
        if isinstance(value, Decimal):
            value = format(value, 'f')
        cell = WriteOnlyCell(sheet, value=value)
        if isinstance(value, str):
            cell.data_type = 's'
            cell.number_format = '@'
        cells.append(cell)
    sheet.append(cells)


def _close_workbook(book):
    # openpyxl 3.1 write-only worksheets own temporary XML separately from
    # Workbook.close(). Keep this adapter-only cleanup off the public contract.
    for sheet in book.worksheets:
        writer = sheet._writer
        if writer is not None and Path(writer.out).exists():
            try:
                if not sheet.closed:
                    sheet.close()
            finally:
                try:
                    writer.close()
                finally:
                    writer.cleanup()
    book.close()


def generate(engine, generated_at):
    output = SpooledTemporaryFile(max_size=1024 * 1024, mode='w+b')
    book = Workbook(write_only=True)
    try:
        sheet = book.create_sheet('Sites')
        _row(sheet, [key for key, _ in COLUMNS])
        count = 0
        with engine.connect().execution_options(isolation_level='REPEATABLE READ') as connection:
            with connection.begin():
                result = connection.execute(QUERY.execution_options(stream_results=True, yield_per=500))
                try:
                    for record in result.mappings():
                        if count >= 1_048_575:
                            raise ValueError('site_reference_exceeds_worksheet_capacity')
                        values = dict(record)
                        values['site_id'] = str(values['site_id'])
                        values['customer_id'] = str(values['customer_id'])
                        values['product_lines'] = ';'.join(line for line, present in
                            (('PV', values['has_pv']), ('ESS', values['has_ess'])) if present) or None
                        values['commissioning_date'] = values['commissioning_date'].isoformat() if values['commissioning_date'] else None
                        _row(sheet, [values[key] for key, _ in COLUMNS])
                        count += 1
                finally:
                    result.close()
        info = book.create_sheet('Instructions')
        for values in (
            ('schema_version', SCHEMA_VERSION),
            ('generated_at', generated_at.isoformat()),
            ('scope', 'all_readable_sites'),
            ('site_count', count),
            ('说明', '全部有权读取的当前现场，包含无设备现场；不受页面筛选、分页或产品线切换影响。'),
            ('空值', '未提供保持空白，不表示0；此文件仅作引用准备，不用于覆盖既有现场。'),
            *COLUMNS,
        ):
            _row(info, values)
        # openpyxl's save_workbook does not close its ZIP on write failure.
        # Own that resource here so it closes before the underlying output.
        with ZipFile(output, 'w', ZIP_DEFLATED, allowZip64=True) as archive:
            ExcelWriter(book, archive).save()
        output.seek(0)
        return SiteExportFile(output, f'site-reference-{generated_at:%Y%m%dT%H%M%SZ}.xlsx', generated_at, count, uuid4())
    except Exception as error:
        output.close()
        raise SiteExportUnavailable() from error
    finally:
        _close_workbook(book)

