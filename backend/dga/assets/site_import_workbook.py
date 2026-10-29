"""Asset-owned, bounded v2 site workbook format; never reinterpret v1 references."""
from io import BytesIO
from zipfile import ZipFile
from datetime import date, datetime
from math import isfinite
from defusedxml.ElementTree import fromstring
from openpyxl import Workbook, load_workbook
from openpyxl.utils.cell import coordinate_from_string,column_index_from_string
from .import_workbook import WorkbookError

MAX_BYTES = 10 * 1024 * 1024
COLUMNS = {
    'Customers': ('record_key','decision','customer_id','customer_name','source_note','reason'),
    'Sites': ('record_key','decision','site_id','site_name','customer_ref','location_text','product_lines',
              'pv_power_mw','ess_power_mw','ess_energy_mwh','commissioning_date','grid_status','operation_status','source_note','reason'),
}

def problem(message, sheet='Format', row=0, field='file', code='invalid_site_workbook'):
    return dict(code=code,severity='ERROR',sheet=sheet,row=row,field=field,message=message)

def reject(message, **context): raise WorkbookError([problem(message,**context)])

def template():
    book=Workbook(); book.active.title='Format'
    book.active.append(['version','mode']); book.active.append([2,'SITES'])
    for name,columns in COLUMNS.items():
        sheet=book.create_sheet(name); sheet.append(columns); sheet.freeze_panes='A2'
    sheet=book.create_sheet('Instructions')
    for line in (
        'v2 SITES：仅新增/明确引用客户与现场；不覆盖已有资料。客户表可留空，现场表至少一条。',
        'record_key为表内不同文本；decision填写NEW或REFERENCE。新增必须填写reason。',
        'REFERENCE填写正式UUID；现场customer_ref为已有客户UUID或NEW:客户记录键。',
        '同名仅作候选，核对后明确引用或确认不同现场新增；不能通过导入更改客户归属。',
        'product_lines使用PV、ESS或PV;ESS。PV功率MW；ESS功率MW及容量MWh。未知留空。',
        '日期YYYY-MM-DD；状态用现有枚举。所有编号用文本；不执行公式/宏。',
        'grid_status: NOT_STARTED / IN_PROGRESS / COMPLETED；operation_status: NOT_OPERATIONAL / PARTIAL / OPERATIONAL / DECOMMISSIONED。',
        'source_note可空，最多2000字符；各表最多500行，XLSX最多10MiB，展开50MiB/200成员/150000单元格。',
    ): sheet.append([line])
    output=BytesIO(); book.save(output); return output.getvalue()

def parse(content):
    book=None
    try:
        with ZipFile(BytesIO(content)) as archive:
            entries=archive.infolist()
            if len(entries)>200 or len({e.filename for e in entries})!=len(entries) or sum(e.file_size for e in entries)>50*1024*1024:
                reject('展开超过50 MiB、200成员或存在重复成员；请分批。')
            cells=0
            for entry in entries:
                extent_row=extent_column=0
                parts=entry.filename.replace('\\','/').split('/')
                if entry.flag_bits&1 or '..' in parts or entry.filename.startswith('/') or 'externallinks/' in entry.filename.lower() or entry.filename.lower().endswith('.bin'):
                    reject('拒绝加密、路径、外部链接及宏文件。')
                if entry.filename.endswith(('.xml','.rels')):
                    for node in fromstring(archive.read(entry)).iter():
                        tag=node.tag.rsplit('}',1)[-1]
                        if tag=='c':
                            cells+=1
                            column,row_index=coordinate_from_string(node.get('r','A1'))
                            extent_row=max(extent_row,row_index); extent_column=max(extent_column,column_index_from_string(column))
                            if row_index*column_index_from_string(column)>150000:
                                reject('稀疏坐标会展开超过150000单元格；请移除远处内容/格式并压紧业务行。')
                        if cells>150000: reject('超过150000单元格；请分批。')
                        if tag=='f' or (tag=='Relationship' and node.get('TargetMode')=='External'):
                            reject('拒绝公式及外部链接，请粘贴为值。')
                if extent_row*extent_column>150000: reject('工作表稀疏矩形展开超过150000单元格；请压紧业务行列。')
        book=load_workbook(BytesIO(content),read_only=True,data_only=False,keep_links=False)
        if set(book.sheetnames)-{'Format','Instructions','Customers','Sites','Materials','Assets'} or 'Format' not in book or 'Sites' not in book:
            reject('请使用v2现场模板，必须包含Format及Sites。')
        book['Format'].reset_dimensions()
        marker=[]
        for values in book['Format'].iter_rows(values_only=True):
            values=list(values)
            while values and values[-1] is None: values.pop()
            if values: marker.append(tuple(values))
        if marker != [('version','mode'),(2,'SITES')]:
            reject('Format必须只包含version=2、mode=SITES，不接收额外字段。')
        result=[]
        for name in ('Customers','Sites','Materials','Assets'):
            if name not in book: continue
            sheet=book[name]; sheet.reset_dimensions()
            iterator=sheet.iter_rows()
            header=next(iterator,())
            headers=[cell.value for cell in header]
            while headers and headers[-1] is None: headers.pop()
            if name in ('Assets','Materials'):
                if any(any(cell.value is not None for cell in row) for row in iterator): reject('仅现场模式不接收设备/物料业务行。',sheet=name)
                continue
            if not headers or any(not isinstance(h,str) for h in headers) or len(set(headers))!=len(headers) or set(headers)-set(COLUMNS[name]) or 'record_key' not in headers:
                reject('列名缺失、重复或不受支持；请重新下载v2模板。',sheet=name,row=1,field='headers')
            count=0
            for n,row in enumerate(iterator,2):
                if not any(c.value is not None for c in row): continue
                count+=1
                if count>500: reject('每表最多500条业务行；请分批。',sheet=name,row=n)
                if any(c.value is not None for c in row[len(headers):]): reject('数据列缺少列名。',sheet=name,row=n)
                values={}
                for key,cell in zip(headers,row):
                    value=cell.value
                    if isinstance(value,str) and (len(value)>2000 or any(ord(ch)<32 and ch not in '\n\r\t' for ch in value)):
                        reject('文本过长或含无效字符；来源备注最多2000字符。',sheet=name,row=n,field=key)
                    if isinstance(value,float) and not isfinite(value): reject('必须填写有限数值。',sheet=name,row=n,field=key)
                    if isinstance(value,(date,datetime)): value=value.isoformat()
                    values[key]=value
                result.append(dict(sheet=name,row=n,values=values))
        if not any(r['sheet']=='Sites' for r in result): reject('Sites没有现场业务行；仅说明/客户不能发布。',sheet='Sites')
        return result
    except WorkbookError: raise
    except Exception as error: raise WorkbookError([problem('无法安全读取v2现场XLSX；请使用模板并粘贴数据值。')]) from error
    finally:
        if book: book.close()
