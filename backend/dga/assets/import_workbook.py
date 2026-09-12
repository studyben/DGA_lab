"""Private bounded XLSX adapter; business validation belongs to assets."""
from io import BytesIO
from datetime import date, datetime
from zipfile import ZipFile
from defusedxml.ElementTree import fromstring
from openpyxl import load_workbook
from openpyxl.utils.cell import coordinate_from_string, column_index_from_string
from .import_validation import REQUIRED, OPTIONAL, issue

MAX_BYTES = 2 * 1024 * 1024


class WorkbookError(Exception):
    def __init__(self, issues):
        self.issues = issues


def reject(message, row=0, field='file'):
    raise WorkbookError([issue(row, field, message)])


def inspect_archive(content):
    with ZipFile(BytesIO(content)) as archive:
        entries = archive.infolist()
        if len(entries) > 100 or len({e.filename for e in entries}) != len(entries) or sum(e.file_size for e in entries) > 10 * 1024 * 1024:
            reject('文件展开后超过10 MiB、成员重复或超过100个成员；请拆分或重新保存模板。')
        cells = 0
        for entry in entries:
            if entry.flag_bits & 1 or 'externallinks/' in entry.filename.lower() or entry.filename.lower().endswith('.bin'):
                reject('不接受加密、外部链接或宏；请使用普通 XLSX 模板。')
            data = archive.read(entry)
            if entry.filename.endswith(('.xml', '.rels')):
                root = fromstring(data)
                for node in root.iter():
                    tag = node.tag.rsplit('}', 1)[-1]
                    if tag == 'Relationship' and node.get('TargetMode') == 'External':
                        reject('文件包含外部链接；请移除后上传。')
                    if tag == 'c':
                        column, _ = coordinate_from_string(node.get('r', 'A1'))
                        if column_index_from_string(column) > 32:
                            reject('第32列以外存在内容；请移除模板外数据后上传。')
                        cells += 1
                        if cells > 10000:
                            reject('文件超过10000个单元格；请拆分批次。')
                    if tag == 'row' and int(node.get('r', '0')) > 501:
                        reject('模板最多500条资产行；请移除远处的空白格式或拆分批次。')
                    if tag == 'f':
                        reject('不接受公式；请粘贴为值后重新上传。')


def parse(content):
    book = None
    try:
        inspect_archive(content)
        book = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
        if 'Assets' not in book.sheetnames or set(book.sheetnames) - {'Assets', 'Sites', 'Materials', 'Instructions'}:
            reject('请使用模板中的 Assets 表，不接受其他业务数据表。')
        sheet = book['Assets']
        # Do not trust user-controlled worksheet dimension hints.
        sheet.reset_dimensions()
        iterator = sheet.iter_rows(max_row=501, max_col=32)
        headers = [cell.value for cell in next(iterator)]
        while headers and headers[-1] is None:
            headers.pop()
        if any(not isinstance(h, str) for h in headers) or len(set(headers)) != len(headers):
            reject('列名为空或重复；请重新下载模板。', 1, 'headers')
        missing = set(REQUIRED) - set(headers)
        if missing:
            reject('缺少必填列：' + ', '.join(sorted(missing)), 1, sorted(missing)[0])
        rows = []
        for n, cells in enumerate(iterator, 2):
            if all(cell.value is None for cell in cells):
                continue
            if any(cell.value is not None for cell in cells[len(headers):]):
                reject('有数据的列必须有列名。', n, 'headers')
            values = {}
            for key, cell in zip(headers, cells):
                value = cell.value
                if isinstance(value, str) and (len(value) > 1000 or '\x00' in value):
                    reject('单元格过长或包含无效字符；最多1000字符。', n, key)
                if isinstance(value, (datetime, date)):
                    # Date cells are not guessed into UTC timestamps; validator rejects timezone-free time.
                    value = value.isoformat()
                values[key] = value
            rows.append({'row': n, 'values': values})
        if not rows:
            reject('Assets 表没有资产数据。')
        return rows
    except WorkbookError:
        raise
    except Exception as error:
        raise WorkbookError([issue(0, 'file', '无法安全读取 XLSX 文件；请重新下载模板并粘贴数据值。')]) from error
    finally:
        if book:
            book.close()
