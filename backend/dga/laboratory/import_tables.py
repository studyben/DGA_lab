"""Laboratory format1 template OUTPUT only. Parsing belongs to a later approved slice."""
import csv
import json
from io import BytesIO, StringIO
from zipfile import ZipFile, ZIP_DEFLATED
from openpyxl import Workbook
from .import_contracts import XLSX_TYPE

SAMPLE_COLUMNS = (
    'source_lab_key','source_lab_name','external_sample_id','source_site_name','source_whole_unit_serial',
    'source_transformer_serial','formal_asset_id','system_asset_number','sampled_at','sampled_timezone',
    'sampling_context','physical_received','received_at','received_timezone','container_count',
    'report_filenames','notes','source_note')
TEST_COLUMNS = (
    'source_lab_key','external_sample_id','external_test_id','test_type','measured_at','measured_timezone',
    'source_method','source_method_version','method_version_id','report_selected','source_instrument',
    'source_analyst','source_qa','h2_value','h2_qualifier','h2_unit','ch4_value','ch4_qualifier','ch4_unit',
    'c2h2_value','c2h2_qualifier','c2h2_unit','c2h4_value','c2h4_qualifier','c2h4_unit',
    'c2h6_value','c2h6_qualifier','c2h6_unit','co_value','co_qualifier','co_unit',
    'co2_value','co2_qualifier','co2_unit','result_value','result_qualifier','result_unit','notes','source_note')

EXPLANATIONS = {
    'source_lab_key':'批次内来源别名：字母开头、字母数字下划线/短横线，最多80；须显式新增或引用稳定来源身份。',
    'source_lab_name':'来源实验室原名，最多200；名称不自动匹配或新建来源。',
    'external_sample_id':'来源油样标识，文本1–160；保留大小写和前导零，检测表使用同一来源别名及标识。',
    'external_test_id':'来源检测标识，文本1–160；每次完整检测或复测不同，不使用数值/日期替代。',
    'source_site_name':'来源现场陈述，最多200；不代表历史归属已核实。',
    'source_whole_unit_serial':'来源整机序列号，文本最多160；不自动绑定正式资产。',
    'source_transformer_serial':'来源变压器序列号，文本最多160；不按序列号与日期合并。',
    'formal_asset_id':'可选正式物理变压器UUID；未知留空，不创建临时资产。',
    'system_asset_number':'可选正式系统资产编号，最多80；不替代油样外部标识。',
    'sampled_at':'采样原时间文本<=80：YYYY-MM-DD或ISO完整时刻；DATE不补午夜。',
    'measured_at':'测量原时间文本<=80：YYYY-MM-DD或ISO完整时刻；缺失保持空白。',
    'received_at':'真实收到实物时的完整ISO时刻；未收实物留空，不能填上传时间。',
    'sampled_timezone':'明确IANA来源时区<=100；不默认浏览器/系统时区。',
    'measured_timezone':'明确IANA来源时区<=100；夏令时重复小时须显式偏移。',
    'received_timezone':'真实收样原时区<=100；无实物不补造。',
    'sampling_context':'SITE / REPAIR_CENTER / OTHER；未知留空待核对。',
    'physical_received':'TRUE/FALSE，明确是否收到实物；不从日期自动判断。',
    'container_count':'未收实物空或0；收到则真实1–20；暂存不创建容器。',
    'report_filenames':'可空或JSON数组（最多100文件名，每名255）；拒绝路径，仅声明PDF，文件始终未提供。',
    'test_type':'DGA / MOISTURE / BREAKDOWN_VOLTAGE；未知类型仍须保留原行。',
    'source_method':'原始方法文本<=160；不猜ASTM兼容性。',
    'source_method_version':'来源方法版本原文<=80；不自动关联系统版本。',
    'method_version_id':'明确批准配置的方法UUID；未知留空，不使用MVP占位方法。',
    'report_selected':'可空或TRUE/FALSE；多条同类检测须明确选择，不自动平均/取最新。',
    'source_instrument':'第三方原仪器文本<=160；不替换为本地仪器或校准通过。',
    'source_analyst':'原分析员文本<=150；未知留空，不替换为上传/核对人。',
    'source_qa':'原质量资料文本<=2000；缺失不推导PASS。',
    'notes':'普通业务备注，可空<=2000字符；与来源备注分别保存。',
    'source_note':'本行资料依据，可空<=2000字符；不作为身份或重复键。'
}


def instructions(example):
    fields = dict(EXPLANATIONS)
    for code in TEST_COLUMNS:
        if code.endswith('_value'):
            fields[code] = '十进制原值；0有效，空白不是ND；全局小于1e12/最多6位小数，仍须批准方法的更严格精度。'
        elif code.endswith('_qualifier'):
            fields[code] = 'EQ / ND / LT / GT；ND无数值，EQ/LT/GT须有有限非负值/边界。'
        elif code.endswith('_unit'):
            fields[code] = '原单位文本<=40；必须与批准方法字段一致，不做转换或兼容推断。'
    return dict(version=1, format='LABORATORY', example_is_synthetic=example,
        columns=dict(samples=list(SAMPLE_COLUMNS), tests=list(TEST_COLUMNS)), fields=fields,
        rules=[
            '仅暂存资料，不创建正式油样、检测、容器或有效结果；PDF仅声明未提供。',
            'DGA每行完整H2/CH4/C2H2/C2H4/C2H6/CO/CO2；微水/击穿用result三列，复测另行。',
            '两份CSV同批UTF-8提交；所有表头必须完整，不能映射任意第三方列。',
            '500油样/2000完整检测；XLSX10MiB；每CSV10MiB、两份20MiB；展开50MiB/200成员/150000单元格。',
            '编号和日期使用文本，保留大小写及前导零；不执行公式、宏或外部链接。',
            '原件私有保留；批次用途HISTORICAL/DAILY、来源身份和来源时区在页面显式确认。',
            '示例为虚构资料；资产、批准方法及单位故意未确认，不能当作科学配置或正式有效结果。'
        ])


def example_rows():
    sample = dict(source_lab_key='LAB_A', source_lab_name='虚构示例实验室', external_sample_id='00017',
                  source_site_name='示例现场（未核实）', sampled_at='2026-01-15',
                  sampling_context='OTHER', physical_received='FALSE', container_count='0',
                  notes='虚构示例，含逗号, 引号"与换行\n请核对原始资料', source_note='示例资料，不是正式证据')
    tests = []
    for index in (1, 2):
        row = dict(source_lab_key='LAB_A', external_sample_id='00017', external_test_id=f'DGA-0{index}',
                   test_type='DGA', measured_at='2026-01-16', source_note='独立复测示例；方法与单位待确认')
        for code in ('h2','ch4','c2h2','c2h4','c2h6','co','co2'):
            row[code+'_value'], row[code+'_qualifier'] = '1.25', 'EQ'
        row['h2_value'], row['h2_qualifier'] = '0', 'EQ'
        row['ch4_value'], row['ch4_qualifier'] = '', 'ND'
        tests.append(row)
    return [sample], tests


def build_template(format, example=False):
    info = instructions(example)
    samples, tests = example_rows() if example else ([], [])
    if format == 'CSV_PAIR':
        output = BytesIO()
        with ZipFile(output, 'w', ZIP_DEFLATED) as archive:
            for name, columns, rows in (('samples.csv', SAMPLE_COLUMNS, samples), ('tests.csv', TEST_COLUMNS, tests)):
                stream = StringIO(newline='')
                writer = csv.writer(stream)
                writer.writerow(columns)
                writer.writerows([[row.get(c, '') for c in columns] for row in rows])
                archive.writestr(name, stream.getvalue().encode('utf-8'))
            archive.writestr('instructions.json', json.dumps(info, ensure_ascii=False, indent=2).encode('utf-8'))
        return dict(content=output.getvalue(), content_type='application/zip',
                    filename='laboratory-import-v1'+('-example' if example else '')+'-csv.zip')
    book = Workbook()
    book.active.title = 'Format'
    book.active.append(['version', 'format'])
    book.active.append([1, 'LABORATORY'])
    for name, columns, rows in (('油样信息', SAMPLE_COLUMNS, samples), ('检测结果', TEST_COLUMNS, tests)):
        sheet = book.create_sheet(name)
        sheet.append(columns)
        for row in rows:
            sheet.append([row.get(c, '') for c in columns])
        sheet.freeze_panes = 'A2'
        for row in sheet:
            for cell in row:
                cell.data_type, cell.number_format = 's', '@'
    sheet = book.create_sheet('Instructions')
    for rule in info['rules']:
        sheet.append([rule])
    for code, explanation in info['fields'].items():
        sheet.append([code, explanation])
    for row in sheet:
        for cell in row:
            if cell.value is not None:
                cell.data_type = 's'
    result = BytesIO()
    book.save(result)
    return dict(content=result.getvalue(), content_type=XLSX_TYPE,
                filename='laboratory-import-v1'+('-example' if example else '')+'.xlsx')

