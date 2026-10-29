"""v2 equipment staging behavior through AssetImports and real PostgreSQL."""
from datetime import datetime, timezone
from io import BytesIO
from uuid import uuid4

import pytest
from openpyxl import Workbook, load_workbook
from dga.assets.public import AssetImports, AssetDirectory, AssetLifecycle, AssetQueryError
from tests.test_asset_import import import_context
from tests.test_site_import import Files
from tests.test_sample_reception import SITE_ID

HEADERS = {
    'Customers': 'record_key decision customer_id customer_name source_note reason'.split(),
    'Sites': 'record_key decision site_id site_name customer_ref location_text product_lines pv_power_mw ess_power_mw ess_energy_mwh commissioning_date grid_status operation_status source_note reason'.split(),
    'Materials': 'record_key decision material_number model asset_type source_note reason'.split(),
    'Assets': 'record_key decision serial_number material_ref product_line machine_type status location_kind effective_at parent_record_key site_id site_name power_mw energy_mwh equipment_name tag_number commissioning_date battery_manufacturer source_note reason'.split(),
}


def document(assets=None, *, materials=(), sites=(), customers=(), mode='ASSETS'):
    book = Workbook()
    book.active.title = 'Format'
    book.active.append(['version', 'mode'])
    book.active.append([2, mode])
    for name, rows in [('Customers', customers), ('Sites', sites), ('Materials', materials),
                       ('Assets', assets if assets is not None else [equipment()])]:
        sheet = book.create_sheet(name)
        columns = HEADERS[name] + sorted({key for row in rows for key in row} - set(HEADERS[name]))
        sheet.append(columns)
        for row in rows:
            sheet.append([row.get(key) for key in columns])
            for cell in sheet[sheet.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = 's'
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def equipment(**changes):
    return dict(record_key='root', decision='NEW', serial_number='00001',
                material_ref='MATERIAL:IMPORT-UNIT', product_line='PV', machine_type='INVERTER_UNIT',
                status='IN_SERVICE', location_kind='SITE', effective_at='2025-01-01T08:00:00Z',
                site_id=str(SITE_ID), site_name='Prairie Sun', source_note='Shipping evidence',
                reason='Verified a new physical device') | changes


def context(database_url):
    engine, actor, _, _ = import_context(database_url)
    files = Files()
    imports = AssetImports(engine, files, clock=lambda: datetime(2026, 1, 1, tzinfo=timezone.utc))
    return engine, actor, imports, files


def upload(imports, actor, content=None, **kwargs):
    return imports.submit(actor, filename='devices.xlsx', content=content or document(),
                          format_version=2, mode='ASSETS', idempotency_key=str(uuid4()), **kwargs)


def ready(imports, actor, content=None):
    batch = upload(imports, actor, content)
    assert imports.validate_next()
    return imports.get(actor, batch['id'])


def test_v2_assets_template_and_staging_preserve_original_without_business_writes(database_url):
    engine, actor, imports, _ = context(database_url)
    before = AssetLifecycle(engine).catalog(actor)['total']
    template = load_workbook(BytesIO(imports.template(actor)))
    assert list(template['Format'].values) == [('version', 'mode'), (2, 'ASSETS')]
    assert list(next(template['Assets'].values)) == HEADERS['Assets']
    content = document()
    batch = upload(imports, actor, content)
    assert batch['state'] == 'STAGED'
    assert batch['operation_result']['outcome'] == 'STAGED'
    assert imports.validate_next()
    preview = imports.get(actor, batch['id'])
    assert preview['state'] == 'VALIDATED'
    assert preview['summary']['error_count'] == 0, preview['issues']
    assert preview['rows'][0]['references']['site'] == {'kind': 'EXISTING', 'site_id': str(SITE_ID)}
    assert preview['rows'][0]['values']['serial_number'] == '00001'
    assert imports.source_file(actor, batch['id'])['content'] == content
    with pytest.raises(AssetQueryError, match='import_mode_not_publishable'):
        imports.publish(actor, batch['id'], validation_revision=preview['validation_revision'],
                        idempotency_key=str(uuid4()), acknowledged_warning_ids=[])
    assert AssetLifecycle(engine).catalog(actor)['total'] == before
    engine.dispose()


def test_explicit_existing_and_batch_references_resolve_independent_of_row_order(database_url):
    from tests.test_sample_reception import CUSTOMER_ID
    engine, actor, imports, _ = context(database_url)
    content = document([
        equipment(record_key='child', serial_number='child-01', material_ref='NEW:existing-material',
                  machine_type='TRANSFORMER', location_kind='PARENT', parent_record_key='root',
                  site_id=None, site_name=None, effective_at='2025-01-01T09:00:00Z'),
        equipment(material_ref='NEW:new-material', site_id='NEW:new-site', site_name='New project'),
        equipment(record_key='existing-root', serial_number='existing-root-01'),
    ], materials=[
        dict(record_key='new-material', decision='NEW', material_number='NEW:UNIT/01', model='New approved model',
             asset_type='WHOLE_UNIT', source_note='Catalog authorization', reason='Verified distinct specification'),
        dict(record_key='existing-material', decision='REFERENCE', material_number='IMPORT-TX', model='TX-4400',
             asset_type='TRANSFORMER', reason='Reference approved catalog'),
    ], customers=[dict(record_key='customer', decision='NEW', customer_name='New owner', reason='Verified organization')],
       sites=[dict(record_key='new-site', decision='NEW', site_name='New project', customer_ref='NEW:customer',
                   product_lines='PV', reason='Verified independent project')])
    before = AssetLifecycle(engine).catalog(actor)['total']
    preview = ready(imports, actor, content)
    assert preview['summary']['error_count'] == 0, preview['issues']
    assert preview['summary']['total_rows'] == 7
    rows = {r['values']['record_key']: r for r in preview['rows']}
    assert rows['root']['references']['site'] == {'kind': 'BATCH', 'record_key': 'new-site'}
    assert rows['root']['values']['model'] == 'New approved model'
    assert rows['child']['values']['material_number'] == 'IMPORT-TX'
    assert rows['child']['resolved_context']['site_name'] == 'New project'
    assert rows['child']['references']['site'] == rows['root']['references']['site']
    assert rows['existing-root']['resolved_context']['customer_id'] == str(CUSTOMER_ID)
    assert rows['new-material']['values']['source_note'] == 'Catalog authorization'
    assert AssetLifecycle(engine).catalog(actor)['total'] == before
    engine.dispose()


@pytest.mark.parametrize('variant,field,code', [
    ('new-existing', 'material_number', 'material_exists_choose_reference'),
    ('model-conflict', 'model', 'material_conflict'),
    ('type-conflict', 'asset_type', 'material_conflict'),
    ('duplicate-number', 'material_number', 'duplicate_material_number'),
    ('unknown-type', 'asset_type', 'invalid_value'),
    ('missing-decision', 'decision', 'invalid_value'),
    ('forbidden-site-write', 'location_text', 'invalid_column'),
    ('forbidden-customer', 'customer_id', 'invalid_column'),
])
def test_material_conflicts_and_forbidden_site_writes_are_actionable(database_url, variant, field, code):
    engine, actor, imports, _ = context(database_url)
    material = dict(record_key='m', decision='REFERENCE', material_number='IMPORT-UNIT', model='SG4400UD',
                    asset_type='WHOLE_UNIT', reason='Checked catalog')
    rows = [equipment()]
    materials = [material]
    if variant == 'new-existing': material['decision'] = 'NEW'
    elif variant == 'model-conflict': material['model'] = 'Conflicting model'
    elif variant == 'type-conflict': material['asset_type'] = 'TRANSFORMER'
    elif variant == 'duplicate-number': materials.append(material | dict(record_key='m2'))
    elif variant == 'unknown-type': material.update(decision='NEW', material_number='UNKNOWN-TYPE', asset_type='PCS')
    elif variant == 'missing-decision': material['decision'] = None
    elif variant == 'forbidden-site-write': rows[0]['location_text'] = 'Cannot write site from device'
    else: rows[0]['customer_id'] = str(uuid4())
    content = document(rows, materials=materials)
    batch = upload(imports, actor, content)
    if batch['state'] == 'STAGED':
        imports.validate_next()
        batch = imports.get(actor, batch['id'])
    assert any(i['field'] == field and i['code'] == code and i['severity'] == 'ERROR' for i in batch['issues']), batch['issues']
    assert imports.source_file(actor, batch['id'])['content'] == content
    assert AssetLifecycle(engine).catalog(actor)['total'] == 2
    engine.dispose()


@pytest.mark.parametrize('change,field', [
    ({'serial_number': 123}, 'serial_number'), ({'serial_number': None}, 'serial_number'),
    ({'decision': 'REFERENCE'}, 'decision'), ({'reason': ' '}, 'reason'),
    ({'status': 'SPARE'}, 'status'), ({'location_kind': 'WAREHOUSE'}, 'location_kind'),
    ({'effective_at': '2025-01-01'}, 'effective_at'),
    ({'effective_at': '2027-01-01T00:00:00Z'}, 'effective_at'),
    ({'product_line': 'ESS', 'machine_type': 'ESS_SYSTEM'}, 'site_id'),
    ({'machine_type': 'TRANSFORMER'}, 'machine_type'),
    ({'power_mw': -1}, 'power_mw'), ({'power_mw': '0.0000001'}, 'power_mw'),
    ({'power_mw': True}, 'power_mw'), ({'energy_mwh': 1}, 'energy_mwh'),
    ({'commissioning_date': 'yesterday'}, 'commissioning_date'),
    ({'site_id': str(uuid4())}, 'site_id'), ({'site_id': 'NEW:absent'}, 'site_id'),
    ({'material_ref': 'MATERIAL:absent'}, 'material_ref'),
])
def test_equipment_scalar_position_and_time_failures_are_actionable(database_url, change, field):
    engine, actor, imports, _ = context(database_url)
    preview = ready(imports, actor, document([equipment(**change)]))
    problems = [i for i in preview['issues'] if i['severity'] == 'ERROR' and i['field'] == field]
    assert problems, preview['issues']
    assert problems[0]['original_value'] == equipment(**change)[field]
    assert problems[0]['suggestion']
    engine.dispose()


def test_site_id_name_difference_requires_explicit_id_review(database_url):
    engine, actor, imports, _ = context(database_url)
    preview = ready(imports, actor, document([equipment(site_name='Source says another name')]))
    assert preview['summary']['error_count'] == 0
    warning = next(i for i in preview['issues'] if i['code'] == 'site_name_difference')
    assert warning['severity'] == 'WARNING'
    assert len(warning['id']) == 24
    row = preview['rows'][0]
    assert row['references']['site']['site_id'] == str(SITE_ID)
    assert row['resolved_context']['site_name'] == 'Prairie Sun'
    assert row['original_values']['site_name'] == 'Source says another name'
    name_only = ready(imports, actor, document([equipment(site_id=None)]))
    assert any(i['code'] == 'site_reference_required' for i in name_only['issues'])
    assert str(SITE_ID) in [c['id'] for c in name_only['rows'][0]['candidates']]
    engine.dispose()



@pytest.mark.parametrize('variant,field,code', [
    ('missing', 'parent_record_key', 'missing_parent'),
    ('self', 'parent_record_key', 'cyclic_installation'),
    ('cycle', 'parent_record_key', 'cyclic_installation'),
    ('duplicate-key', 'record_key', 'invalid_value'),
    ('early-child', 'effective_at', 'invalid_effective_at'),
    ('line', 'parent_record_key', 'invalid_value'),
    ('repair-parent', 'parent_record_key', 'parent_not_at_site'),
    ('child-site', 'site_id', 'inherited_site_conflict'),
    ('name-only-child', 'site_name', 'inherited_site_conflict'),
])
def test_initial_tree_position_status_and_time_failures_block(database_url, variant, field, code):
    engine, actor, imports, _ = context(database_url)
    root = equipment()
    child = equipment(record_key='child', serial_number='new-child', material_ref='MATERIAL:IMPORT-TX',
                      machine_type='TRANSFORMER', location_kind='PARENT', parent_record_key='root',
                      site_id=None, site_name=None, effective_at='2025-01-01T09:00:00Z')
    if variant == 'missing': child['parent_record_key'] = str(uuid4())
    elif variant == 'self': child['parent_record_key'] = 'child'
    elif variant == 'cycle': root.update(location_kind='PARENT', parent_record_key='child', site_id=None, site_name=None)
    elif variant == 'duplicate-key': child['record_key'] = 'root'
    elif variant == 'early-child': child['effective_at'] = '2024-01-01T00:00:00Z'
    elif variant == 'line': child['product_line'] = 'ESS'
    elif variant == 'repair-parent': root.update(location_kind='REPAIR_CENTER', site_id=None, site_name=None, status='SPARE')
    elif variant == 'child-site': child['site_id'] = str(uuid4())
    else: child['site_name'] = 'Wrong inherited name'
    preview = ready(imports, actor, document([child, root]))
    assert any(i['field'] == field and i['code'] == code and i['severity'] == 'ERROR' for i in preview['issues']), preview['issues']
    engine.dispose()


def test_100_node_initial_path_is_valid_and_101_is_blocked(database_url):
    engine, actor, imports, _ = context(database_url)
    rows = [equipment(record_key='n0')]
    for n in range(1, 101):
        rows.append(equipment(record_key=f'n{n}', serial_number=f'depth-{n}',
                              location_kind='PARENT', parent_record_key=f'n{n-1}', site_id=None, site_name=None))
    legal = ready(imports, actor, document(list(reversed(rows[:100]))))
    assert legal['summary']['error_count'] == 0, legal['issues']
    assert legal['total_rows'] == 100
    assert legal['rows'][0]['resolved_context']['site_name'] == 'Prairie Sun'
    invalid = ready(imports, actor, document(list(reversed(rows))))
    assert any(i['code'] == 'cyclic_installation' and i['field'] == 'parent_record_key' for i in invalid['issues'])
    engine.dispose()


@pytest.mark.parametrize('status', ['SPARE', 'UNDER_REPAIR', 'RETIRED'])
def test_repair_center_initial_transformer_and_matching_child_site_are_valid(database_url, status):
    engine, actor, imports, _ = context(database_url)
    preview = ready(imports, actor, document([
        equipment(record_key='repair', serial_number='repair-01', material_ref='MATERIAL:IMPORT-TX',
                  machine_type='TRANSFORMER', location_kind='REPAIR_CENTER', status=status, site_id=None, site_name=None),
        equipment(record_key='child', serial_number='child-01', location_kind='PARENT', parent_record_key='root'),
        equipment(),
    ]))
    assert preview['summary']['error_count'] == 0, preview['issues']
    assert preview['rows'][1]['references']['site']['site_id'] == str(SITE_ID)
    assert not preview['rows'][0]['references'].get('site')
    engine.dispose()


def test_duplicate_serial_candidates_identify_existing_and_same_batch_devices(database_url):
    engine, actor, imports, _ = context(database_url)
    preview = ready(imports, actor, document([
        equipment(serial_number='INV-UNIT-7788'),
        equipment(record_key='other', serial_number='inv-unit-7788'),
    ]))
    warnings = [i for i in preview['issues'] if i['code'] == 'duplicate_serial']
    assert len(warnings) == 2 and len({i['id'] for i in warnings}) == 2
    assert preview['summary']['error_count'] == 0
    candidates = preview['rows'][0]['candidates']
    existing = next(c for c in candidates if c['kind'] == 'ASSET')
    assert existing['system_asset_number'] == 'SYS-PV-001'
    assert existing['model'] == 'SG4400UD'
    assert existing['site_name'] == 'Prairie Sun'
    assert existing['customer_name'] == 'Prairie Solar LLC'
    assert existing['location']['kind'] == 'SITE'
    assert any(c['kind'] == 'BATCH_ASSET' and c['record_key'] == 'other' for c in candidates)
    assert all(r['values']['decision'] == 'NEW' and not r.get('target_id') for r in preview['rows'])
    engine.dispose()


def test_duplicate_file_history_is_paged_and_preserves_prior_mode_results(database_url):
    from tests.test_site_import import document as site_document
    engine, actor, imports, _ = context(database_url)
    content = site_document()
    original = imports.submit(actor, filename='published-sites.xlsx', content=content, format_version=2,
                              mode='SITES', idempotency_key=str(uuid4()))
    imports.validate_next()
    original = imports.get(actor, original['id'])
    published = imports.publish(actor, original['id'], validation_revision=original['validation_revision'],
                                idempotency_key=str(uuid4()), acknowledged_warning_ids=[])
    for n in range(22):
        last = upload(imports, actor, content)
    first = imports.get(actor, last['id'])
    second = imports.get(actor, last['id'], duplicate_page=2)
    assert first['duplicate_total'] == 22
    assert len(first['duplicate_batches']) == 20
    assert len(second['duplicate_batches']) == 2
    prior = first['duplicate_batches'][0]
    assert prior['id'] == str(published['id'])
    assert prior['mode'] == 'SITES' and prior['format_version'] == 2
    assert prior['state'] == 'PUBLISHED' and prior['published_at']
    engine.dispose()


def test_revise_preserves_source_and_invalidates_current_preview(database_url):
    from concurrent.futures import ThreadPoolExecutor
    engine, actor, imports, _ = context(database_url)
    content = document([equipment(site_name='Wrong source name')])
    before = ready(imports, actor, content)
    key = str(uuid4())
    changes = [dict(sheet='Assets', row=2, values=dict(site_name='Prairie Sun', source_note='  Revised evidence\nline two  ', reason='Checked authoritative site ID'))]
    def revise():
        return imports.revise(actor, before['id'], expected_revision=before['validation_revision'],
                              changes=changes, reason='Checked authoritative site ID', idempotency_key=key)
    with ThreadPoolExecutor(max_workers=2) as pool:
        results = list(pool.map(lambda _: revise(), range(2)))
    result = results[0]
    assert result['operation_result'] == results[1]['operation_result']
    assert result['operation_result']['revision'] == before['validation_revision'] + 1
    assert result['summary']['warning_count'] == 0
    assert result['rows'][0]['original_values']['site_name'] == 'Wrong source name'
    assert result['rows'][0]['values']['source_note'] == '  Revised evidence\nline two  '
    history = imports.revision_history(actor, before['id'])
    assert history['total'] == 1
    assert history['changes'][0]['before']['site_name'] == 'Wrong source name'
    assert history['changes'][0]['after']['site_name'] == 'Prairie Sun'
    assert imports.source_file(actor, before['id'])['content'] == content
    later = imports.revise(actor, before['id'], expected_revision=result['validation_revision'],
                           changes=[dict(sheet='Assets', row=2, values=dict(source_note='Note-only update'))],
                           reason='More evidence', idempotency_key=str(uuid4()))
    replay = revise()
    assert replay['operation_result'] == result['operation_result']
    assert replay['validation_revision'] == later['validation_revision']
    assert replay['rows'][0]['values']['decision'] == 'NEW'
    assert replay['rows'][0]['references']['site']['site_id'] == str(SITE_ID)
    with pytest.raises(AssetQueryError, match='import_operation_conflict'):
        imports.revise(actor, before['id'], expected_revision=before['validation_revision'], changes=changes,
                       reason='Different intent', idempotency_key=key)
    engine.dispose()


@pytest.mark.parametrize('kind', ['upload', 'revise'])
def test_lost_commit_response_recovers_staging_intent_without_losing_original(database_url, kind):
    from contextlib import contextmanager
    engine, actor, imports, files = context(database_url)
    content = document()
    before = ready(imports, actor, content) if kind == 'revise' else None
    class LostResponse:
        def connect(self): return engine.connect()
        @contextmanager
        def begin(self):
            with engine.begin() as c: yield c
            raise ConnectionError('Response lost after durable commit')
    failed = AssetImports(LostResponse(), files)
    key = str(uuid4())
    kwargs = dict(expected_revision=before['validation_revision'], changes=[dict(sheet='Assets', row=2, values=dict(source_note='Reviewed'))], reason='Checked source', idempotency_key=key) if before else dict(filename='devices.xlsx', content=content, format_version=2, mode='ASSETS', idempotency_key=key)
    with pytest.raises(ConnectionError):
        if before: failed.revise(actor, before['id'], **kwargs)
        else: failed.submit(actor, **kwargs)
    operation = imports.operation(actor, key)
    assert operation['outcome'] == ('REVISED' if before else 'STAGED')
    assert imports.source_file(actor, operation['batch_id'])['content'] == content
    replay = imports.revise(actor, before['id'], **kwargs) if before else imports.submit(actor, **kwargs)
    assert replay['operation_result'] == operation
    assert imports.list_batches(actor)['total'] == 1
    engine.dispose()


def test_current_import_grants_are_rechecked_for_staging_reads_and_changes(database_url):
    from tests.identity_seed import seed_legacy_user
    from tests.test_sample_reception import INITIAL_PASSWORD, CHANGED_PASSWORD
    from dga.shared.auth.public import IdentityService, IdentityError
    engine, admin, imports, _ = context(database_url)
    identity = IdentityService(engine)
    manager_id = seed_legacy_user(database_url, 'device-manager', 'Manager', INITIAL_PASSWORD, ['lab_admin'])
    first = identity.login('device-manager', INITIAL_PASSWORD)
    manager = identity.change_password(first.token, INITIAL_PASSWORD, CHANGED_PASSWORD).actor
    key = str(uuid4())
    batch = imports.submit(manager, filename='manager.xlsx', content=document(), format_version=2, mode='ASSETS', idempotency_key=key)
    assert imports.operation(manager, key)['batch_id'] == str(batch['id'])
    with pytest.raises(AssetQueryError, match='import_operation_not_found'):
        imports.operation(admin, key)
    identity.change_roles(admin, manager_id, add=['field_engineer'], remove=['lab_admin'], expected_revision=0)
    calls = [lambda: imports.get(manager, batch['id']), lambda: imports.list_batches(manager),
             lambda: imports.template(manager), lambda: imports.source_file(manager, batch['id']),
             lambda: imports.operation(manager, key), lambda: upload(imports, manager),
             lambda: imports.revise(manager, batch['id'], expected_revision=0, changes=[dict(sheet='Assets', row=2, values=dict(source_note='x'))], reason='x', idempotency_key=key),
             lambda: imports.publish(manager, batch['id'], validation_revision=1, idempotency_key=key, acknowledged_warning_ids=[])]
    for call in calls:
        with pytest.raises(IdentityError): call()
    assert AssetLifecycle(engine).catalog(admin)['total'] == 2
    engine.dispose()


def test_validation_database_failure_retains_original_and_stops_failed_batch(database_url):
    from sqlalchemy import text
    engine, actor, imports, _ = context(database_url)
    content = document()
    batch = upload(imports, actor, content)
    with engine.begin() as c:
        c.execute(text('ALTER TABLE asset_materials RENAME TO issue50_unavailable_materials'))
    try:
        assert imports.validate_next()
        result = imports.get(actor, batch['id'])
        assert result['state'] == 'FAILED'
        assert result['failure_code'] == 'validation_unavailable'
        assert result['summary']['error_count'] == 1
        assert imports.source_file(actor, batch['id'])['content'] == content
        assert imports.validate_next() is False
    finally:
        with engine.begin() as c:
            c.execute(text('ALTER TABLE issue50_unavailable_materials RENAME TO asset_materials'))
    recovered = imports.revise(actor, batch['id'], expected_revision=result['validation_revision'], changes=[dict(sheet='Assets', row=2, values=dict(source_note='Recovered validation'))], reason='Catalog service restored', idempotency_key=str(uuid4()))
    assert recovered['state'] == 'VALIDATED'
    assert recovered['failure_code'] is None
    assert imports.source_file(actor, batch['id'])['content'] == content
    engine.dispose()


@pytest.mark.parametrize('variant', ['missing-header', 'mode-marker', 'empty-optionals', 'materials-only', 'long-note'])
def test_v2_schema_boundaries_retain_invalid_originals_and_allow_empty_optional_sheets(database_url, variant):
    engine, actor, imports, _ = context(database_url)
    book = load_workbook(BytesIO(document()))
    if variant == 'missing-header': book['Assets'].delete_cols(20)
    elif variant == 'mode-marker': book['Format']['B2'] = 'SITES'
    elif variant == 'empty-optionals':
        for name in ('Customers', 'Sites', 'Materials'):
            del book[name]
            book.create_sheet(name)
    elif variant == 'materials-only':
        del book['Assets']
        book['Materials'].append(['m', 'NEW', 'M-01', 'Approved model', 'WHOLE_UNIT', None, 'Verified'])
    else: book['Assets']['S2'] = 'x' * 2001
    output = BytesIO(); book.save(output); content = output.getvalue()
    batch = upload(imports, actor, content)
    assert batch['state'] == ('STAGED' if variant == 'empty-optionals' else 'FAILED'), batch
    if batch['state'] == 'FAILED': assert batch['issues'][0]['message']
    assert imports.source_file(actor, batch['id'])['content'] == content
    engine.dispose()


def test_issue_ids_distinguish_actionable_fields_and_preview_counts_cover_each_kind(database_url):
    from tests.test_site_import import new_site
    from tests.test_sample_reception import CUSTOMER_ID
    engine, actor, imports, _ = context(database_url)
    preview = ready(imports, actor, document([equipment(serial_number=None, reason=None, location_kind=None)]))
    ids = [i['id'] for i in preview['issues']]
    assert len(ids) == len(set(ids)), preview['issues']
    mixed = ready(imports, actor, document(sites=[new_site()], materials=[dict(record_key='m', decision='REFERENCE', material_number='IMPORT-TX', model='TX-4400', asset_type='TRANSFORMER', reason='Checked')], customers=[dict(record_key='c', decision='REFERENCE', customer_id=str(CUSTOMER_ID), customer_name='Prairie Solar LLC', reason='Checked')]))
    assert mixed['summary']['by_kind']['Assets']['new_rows'] == 1
    assert mixed['summary']['by_kind']['Materials']['reference_rows'] == 1
    assert mixed['summary']['by_kind']['Sites']['new_rows'] == 1
    assert mixed['summary']['by_kind']['Customers']['reference_rows'] == 1
    engine.dispose()


def test_v2_assets_template_explains_explicit_references_and_audit_names_match_mode(database_url):
    from dga.shared.auth.public import IdentityService
    engine, actor, imports, _ = context(database_url)
    book = load_workbook(BytesIO(imports.template(actor)))
    instructions = '\n'.join(str(row[0]) for row in book['Instructions'].values)
    for token in ('ASSETS', 'MATERIAL:', 'NEW:', 'PARENT', 'REPAIR_CENTER', '100'):
        assert token in instructions
    batch = ready(imports, actor)
    imports.revise(actor, batch['id'], expected_revision=batch['validation_revision'],
                   changes=[dict(sheet='Assets', row=2, values=dict(source_note='Reviewed evidence'))],
                   reason='Checked original', idempotency_key=str(uuid4()))
    imports.source_file(actor, batch['id'])
    actions = {event['action_code'] for event in IdentityService(engine).audit_events(actor) if event['entity_id'] == batch['id']}
    assert {'ASSET_IMPORT_SUBMIT', 'ASSET_IMPORT_REVISE', 'ASSET_IMPORT_SOURCE_DOWNLOAD'} <= actions
    engine.dispose()


@pytest.mark.parametrize('version,mode', [(1, 'SITES'), (2, 'UNKNOWN'), (3, 'ASSETS'), (True, 'ASSETS'), (2.0, 'ASSETS')])
def test_version_mode_dispatch_never_reinterprets_an_unsupported_pair(database_url, version, mode):
    engine, actor, imports, _ = context(database_url)
    for operation in (lambda: imports.template(actor, format_version=version, mode=mode),
                      lambda: imports.submit(actor, filename='devices.xlsx', content=document(), format_version=version, mode=mode, idempotency_key=str(uuid4()))):
        with pytest.raises(AssetQueryError, match='invalid_import_format') as error: operation()
        assert error.value.status == 422
    assert imports.list_batches(actor)['total'] == 0
    engine.dispose()


def test_unsupported_persisted_mode_cannot_dispatch_to_site_writer(database_url):
    from sqlalchemy import text
    engine, actor, imports, _ = context(database_url)
    batch = ready(imports, actor)
    with engine.begin() as c:
        c.execute(text("UPDATE asset_import_batches SET mode='UNSUPPORTED' WHERE id=:id"), dict(id=batch['id']))
    calls = [lambda: imports.get(actor, batch['id']),
             lambda: imports.revise(actor, batch['id'], expected_revision=batch['validation_revision'], changes=[dict(sheet='Assets', row=2, values=dict(source_note='x'))], reason='x', idempotency_key=str(uuid4())),
             lambda: imports.publish(actor, batch['id'], validation_revision=batch['validation_revision'], idempotency_key=str(uuid4()), acknowledged_warning_ids=[])]
    for call in calls:
        with pytest.raises(AssetQueryError, match='unsupported_import_format') as error: call()
        assert error.value.status == 503
    assert AssetLifecycle(engine).catalog(actor)['total'] == 2
    engine.dispose()


def test_ambiguous_batch_material_reference_blocks_instead_of_choosing_last_row(database_url):
    engine, actor, imports, _ = context(database_url)
    materials = [dict(record_key='same-key', decision='NEW', material_number=number, model='Approved model', asset_type='WHOLE_UNIT', reason='Checked') for number in ('NEW-M1', 'NEW-M2')]
    preview = ready(imports, actor, document([equipment(material_ref='NEW:same-key')], materials=materials))
    assert any(i['sheet'] == 'Assets' and i['code'] == 'batch_reference_not_found' for i in preview['issues']), preview['issues']
    engine.dispose()


def test_duplicate_candidate_confirmation_changes_with_ancestor_lifecycle_context(database_url):
    from tests.test_sample_reception import WHOLE_UNIT_ID
    engine, actor, imports, _ = context(database_url)
    batch = ready(imports, actor, document([equipment(serial_number='TX-CURRENT-2002', material_ref='MATERIAL:IMPORT-TX', machine_type='TRANSFORMER', location_kind='REPAIR_CENTER', status='SPARE', site_id=None, site_name=None)]))
    first_id = next(i['id'] for i in batch['issues'] if i['code'] == 'duplicate_serial')
    lifecycle = AssetLifecycle(engine)
    revision = lifecycle.history(actor, WHOLE_UNIT_ID)['revision']
    lifecycle.change_status(actor, WHOLE_UNIT_ID, status='RETIRED', effective_at=datetime(2025, 6, 1, tzinfo=timezone.utc), reason='Controlled retirement', expected_revision=revision)
    changed = imports.revise(actor, batch['id'], expected_revision=batch['validation_revision'], changes=[dict(sheet='Assets', row=2, values=dict(source_note='Context reviewed'))], reason='Check current candidates', idempotency_key=str(uuid4()))
    second_id = next(i['id'] for i in changed['issues'] if i['code'] == 'duplicate_serial')
    assert first_id != second_id
    candidate = next(a for a in changed['rows'][0]['candidates'] if a['kind'] == 'ASSET')
    assert candidate['location']['path'][0]['lifecycle_status'] == 'RETIRED'
    engine.dispose()


@pytest.mark.parametrize('child_site', [None, 'NEW:absent', str(SITE_ID)])
def test_unresolved_root_site_stays_actionable_for_implicit_and_explicit_children(database_url, child_site):
    engine, actor, imports, _ = context(database_url)
    content = document([
        equipment(site_id='NEW:absent'),
        equipment(record_key='child', serial_number='new-child', location_kind='PARENT', parent_record_key='root',
                  site_id=child_site, site_name=None),
    ])
    preview = ready(imports, actor, content)
    assert preview['state'] == 'VALIDATED', preview['issues']
    assert preview['failure_code'] is None
    assert any(i['code'] == 'batch_reference_not_found' and i['row'] == 2 for i in preview['issues'])
    assert any(i['code'] == 'inherited_site_conflict' and i['row'] == 3 for i in preview['issues'])
    assert not preview['rows'][0]['references'].get('site')
    assert not preview['rows'][1]['references'].get('site')
    assert not preview['rows'][1].get('resolved_context')
    assert imports.source_file(actor, preview['id'])['content'] == content
    engine.dispose()


def test_progressive_site_revision_preserves_original_and_unresolved_row_problems(database_url):
    engine, actor, imports, _ = context(database_url)
    content = document([equipment(), equipment(record_key='child', serial_number='new-child', location_kind='PARENT', parent_record_key='root')])
    original = ready(imports, actor, content)
    pending = imports.revise(actor, original['id'], expected_revision=original['validation_revision'],
        changes=[dict(sheet='Assets', row=row, values=dict(site_id='NEW:absent')) for row in (2, 3)],
        reason='Source needs site verification', idempotency_key=str(uuid4()))
    assert pending['state'] == 'VALIDATED' and pending['summary']['error_count'] > 0
    assert pending['failure_code'] is None
    root_fixed = imports.revise(actor, original['id'], expected_revision=pending['validation_revision'],
        changes=[dict(sheet='Assets', row=2, values=dict(site_id=str(SITE_ID)))], reason='Root site verified', idempotency_key=str(uuid4()))
    assert any(i['row'] == 3 and i['code'] == 'batch_reference_not_found' for i in root_fixed['issues'])
    complete = imports.revise(actor, original['id'], expected_revision=root_fixed['validation_revision'],
        changes=[dict(sheet='Assets', row=3, values=dict(site_id=str(SITE_ID)))], reason='Child agrees with root', idempotency_key=str(uuid4()))
    assert complete['state'] == 'VALIDATED' and complete['summary']['error_count'] == 0, complete['issues']
    assert all(row['original_values']['site_id'] == str(SITE_ID) for row in complete['rows'])
    assert complete['rows'][1]['references']['site'] == {'kind': 'EXISTING', 'site_id': str(SITE_ID)}
    assert imports.revision_history(actor, original['id'])['total'] == 4
    assert imports.source_file(actor, original['id'])['content'] == content
    assert AssetLifecycle(engine).catalog(actor)['total'] == 2
    engine.dispose()


@pytest.mark.parametrize('existing,submitted', [
    ('Straße', 'Straße'), ('Straße', 'STRAßE'), ('Straße', 'STRASSE'), ('Straße', 'strasse'),
    ('INV-UNIT-7788', 'INV-UNIT-7788'), ('INV-UNIT-7788', 'inv-unit-7788'),
])
def test_existing_unicode_and_ascii_serials_keep_explicit_physical_candidates(database_url, existing, submitted):
    from sqlalchemy import text
    from tests.test_sample_reception import WHOLE_UNIT_ID
    engine, actor, imports, _ = context(database_url)
    with engine.begin() as c:
        c.execute(text('UPDATE formal_assets SET serial_number=:serial WHERE id=:id'), dict(serial=existing, id=WHOLE_UNIT_ID))
    content = document([equipment(serial_number=submitted)])
    preview = ready(imports, actor, content)
    assert preview['state'] == 'VALIDATED' and preview['summary']['error_count'] == 0
    assert sum(i['code'] == 'duplicate_serial' and i['severity'] == 'WARNING' for i in preview['issues']) == 1
    candidate = next(c for c in preview['rows'][0]['candidates'] if c['kind'] == 'ASSET')
    assert candidate['id'] == str(WHOLE_UNIT_ID) and candidate['serial_number'] == existing
    assert candidate['system_asset_number'] == 'SYS-PV-001' and candidate['site_name'] == 'Prairie Sun'
    assert preview['rows'][0]['original_values']['serial_number'] == submitted
    assert preview['rows'][0]['values']['decision'] == 'NEW' and not preview['rows'][0].get('target_id')
    assert imports.source_file(actor, preview['id'])['content'] == content
    engine.dispose()


def test_same_batch_unicode_serial_variants_are_candidates_without_binding(database_url):
    engine, actor, imports, _ = context(database_url)
    preview = ready(imports, actor, document([equipment(serial_number='Straße'), equipment(record_key='other', serial_number='STRASSE')]))
    assert sum(i['code'] == 'duplicate_serial' for i in preview['issues']) == 2
    assert all(any(c['kind'] == 'BATCH_ASSET' for c in row['candidates']) for row in preview['rows'])
    assert all(not any(c['kind'] == 'ASSET' for c in row['candidates']) for row in preview['rows'])
    assert all(row['values']['decision'] == 'NEW' and not row.get('target_id') for row in preview['rows'])
    engine.dispose()


@pytest.mark.parametrize('version,mode', [(2, 'UNSUPPORTED'), (3, 'ASSETS'), (1, 'SITES')])
def test_worker_rejects_unknown_persisted_pair_without_consuming_staged_batch(database_url, version, mode):
    from sqlalchemy import text
    engine, actor, imports, _ = context(database_url)
    content = document()
    before = upload(imports, actor, content)
    try:
        with engine.begin() as c:
            c.execute(text('UPDATE asset_import_batches SET format_version=:version,mode=:mode WHERE id=:id'), dict(version=version, mode=mode, id=before['id']))
        with pytest.raises(AssetQueryError, match='unsupported_import_format') as error:
            imports.validate_next()
        assert error.value.status == 503
    finally:
        with engine.begin() as c:
            c.execute(text("UPDATE asset_import_batches SET format_version=2,mode='ASSETS' WHERE id=:id"), dict(id=before['id']))
    restored = imports.get(actor, before['id'])
    for field in ('state', 'validation_revision', 'rows', 'issues', 'failure_code', 'summary'):
        assert restored[field] == before[field]
    assert imports.source_file(actor, before['id'])['content'] == content
    assert imports.validate_next()
    completed = imports.get(actor, before['id'])
    assert completed['state'] == 'VALIDATED' and completed['summary']['error_count'] == 0
    assert AssetLifecycle(engine).catalog(actor)['total'] == 2
    engine.dispose()


@pytest.mark.parametrize('sqlstate', ['40001', '40P01'])
def test_worker_database_contention_rolls_back_claim_and_retries_without_failed_state(database_url, sqlstate):
    from sqlalchemy import event
    engine, actor, imports, _ = context(database_url)
    content = document()
    before = upload(imports, actor, content)
    armed = True
    def inject_database_retry(connection, cursor, statement, parameters, execution_context, many):
        nonlocal armed
        if armed and statement.startswith('SELECT * FROM asset_materials'):
            armed = False
            return f"DO $$ BEGIN RAISE EXCEPTION 'Injected isolated worker contention' USING ERRCODE='{sqlstate}'; END $$", parameters
        return statement, parameters
    event.listen(engine, 'before_cursor_execute', inject_database_retry, retval=True)
    try:
        assert imports.validate_next() is True
    finally:
        event.remove(engine, 'before_cursor_execute', inject_database_retry)
    assert armed is False
    retained = imports.get(actor, before['id'])
    assert retained['state'] == 'STAGED' and retained['validation_revision'] == 0
    assert retained['rows'] == [] and retained['failure_code'] is None
    assert imports.source_file(actor, before['id'])['content'] == content
    assert imports.validate_next()
    completed = imports.get(actor, before['id'])
    assert completed['state'] == 'VALIDATED' and completed['summary']['error_count'] == 0
    engine.dispose()
