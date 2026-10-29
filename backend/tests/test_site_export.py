from datetime import datetime, timezone
from decimal import Decimal
from uuid import uuid4

import pytest
from openpyxl import load_workbook
from sqlalchemy import text

from dga.assets.public import AssetDirectory
from tests.test_sample_reception import reception_context, SITE_ID, CUSTOMER_ID

AS_OF = datetime(2026, 10, 29, 12, 0, tzinfo=timezone.utc)


@pytest.fixture
def export_context(database_url):
    engine, identity, actor = reception_context(database_url)
    empty_id = uuid4()
    with engine.begin() as connection:
        connection.execute(text("""UPDATE sites SET grid_status='COMPLETED',operation_status='OPERATIONAL',
            commissioning_date='2020-06-01' WHERE id=:id"""), {'id': SITE_ID})
        connection.execute(text("""INSERT INTO site_product_lines VALUES
            (:id,'PV',100,NULL),(:id,'ESS',20,80)"""), {'id': SITE_ID})
        connection.execute(text("INSERT INTO sites(id,customer_id,site_name) VALUES (:id,:customer,'Empty site')"),
                           {'id': empty_id, 'customer': CUSTOMER_ID})
    yield engine, identity, actor, empty_id
    engine.dispose()


def test_export_contains_each_site_once_with_official_capacities_and_unknowns(export_context):
    engine, _, actor, empty_id = export_context
    exported = AssetDirectory(engine, clock=lambda: AS_OF).export_sites(actor)
    try:
        book = load_workbook(exported.file, read_only=True)
        assert book.sheetnames == ['Sites', 'Instructions']
        rows = list(book['Sites'].iter_rows(max_col=13, values_only=True))
        assert rows[0] == ('site_id', 'site_name', 'customer_id', 'customer_name', 'location_text',
                           'product_lines', 'pv_power_mw', 'ess_power_mw', 'ess_energy_mwh',
                           'commissioning_date', 'grid_status', 'operation_status', 'source_note')
        assert len(rows) == 3
        by_id = {row[0]: dict(zip(rows[0], row)) for row in rows[1:]}
        site = by_id[str(SITE_ID)]
        assert site['product_lines'] == 'PV;ESS'
        assert site['pv_power_mw'] == '100.000000'
        assert site['ess_power_mw'] == '20.000000'
        assert site['ess_energy_mwh'] == '80.000000'
        assert site['commissioning_date'] == '2020-06-01'
        assert site['grid_status'] == 'COMPLETED'
        assert site['operation_status'] == 'OPERATIONAL'
        assert site['source_note'] is None
        assert by_id[str(empty_id)]['pv_power_mw'] is None
        assert by_id[str(empty_id)]['product_lines'] is None
        assert exported.site_count == 2
        assert exported.generated_at == AS_OF
        instructions = dict(book['Instructions'].values)
        assert instructions['schema_version'] == 'site-reference-v1'
        assert instructions['scope'] == 'all_readable_sites'
        assert instructions['site_count'] == 2
        assert instructions['generated_at'] == '2026-10-29T12:00:00+00:00'
        book.close()
    finally:
        exported.file.close()


def test_export_records_scope_and_count_without_claiming_browser_saved(export_context):
    engine, identity, actor, _ = export_context
    exported = AssetDirectory(engine, clock=lambda: AS_OF).export_sites(actor)
    exported.file.close()
    events = [event for event in identity.audit_events(actor) if event['action_code'] == 'SITE_EXPORT_PREPARED']
    assert len(events) == 1
    assert events[0]['actor_user_id'] == actor.user_id
    assert events[0]['entity_id'] == exported.export_id
    assert events[0]['details'] == {
        'scope': 'all_readable_sites', 'site_count': 2, 'schema_version': 'site-reference-v1',
        'generated_at': '2026-10-29T12:00:00+00:00', 'delivery': 'prepared_for_download',
    }

def test_more_than_500_sites_safe_text_exact_decimals_and_identity(export_context):
    engine, _, actor, empty_id = export_context
    with engine.begin() as connection:
        connection.execute(text("UPDATE sites SET site_name=:name,location_text=:location WHERE id=:id"),
                           {'id': SITE_ID, 'name': '=HYPERLINK("https://example.invalid")', 'location': '中文\n第二行'})
        connection.execute(text("UPDATE site_product_lines SET power_mw=1234567890.123456 WHERE site_id=:id AND product_line='PV'"), {'id': SITE_ID})
        connection.execute(text("UPDATE site_product_lines SET power_mw=0 WHERE site_id=:id AND product_line='ESS'"), {'id': SITE_ID})
        connection.execute(text("INSERT INTO sites(id,customer_id,site_name) VALUES (:id,:customer,:name)"),
                           [dict(id=uuid4(), customer=CUSTOMER_ID, name='Same site name') for _ in range(601)])
    exported = AssetDirectory(engine).export_sites(actor)
    try:
        book = load_workbook(exported.file, read_only=True)
        rows = list(book['Sites'].iter_rows(max_col=13))
        assert exported.site_count == 603
        assert len(rows) == 604
        selected = next(row for row in rows[1:] if row[0].value == str(SITE_ID))
        assert selected[1].value == '=HYPERLINK("https://example.invalid")'
        assert selected[1].data_type == 's'
        assert selected[0].data_type == 's'
        assert selected[4].value == '中文\n第二行'
        assert selected[6].value == '1234567890.123456'
        assert selected[7].value == '0.000000'
        assert len({row[0].value for row in rows[1:]}) == 603
        assert 'grid_year' not in [cell.value for cell in rows[0]]
        assert all(cell.data_type != 'f' for row in rows for cell in row)
        book.close()
    finally:
        exported.file.close()


def test_empty_export_has_headers_and_zero_count(export_context):
    engine, _, actor, _ = export_context
    with engine.begin() as connection:
        connection.execute(text('TRUNCATE sites CASCADE'))
    exported = AssetDirectory(engine).export_sites(actor)
    try:
        book = load_workbook(exported.file, read_only=True)
        assert len(list(book['Sites'].values)) == 1
        assert dict(book['Instructions'].values)['site_count'] == 0
        assert exported.site_count == 0
        book.close()
    finally:
        exported.file.close()


def test_read_permission_required_but_no_import_permission(export_context):
    from dataclasses import replace
    from dga.shared.auth.public import IdentityError
    engine, _, actor, _ = export_context
    with pytest.raises(IdentityError, match='permission_denied'):
        AssetDirectory(engine).export_sites(replace(actor, permissions=frozenset()))
    with pytest.raises(IdentityError, match='password_change_required'):
        AssetDirectory(engine).export_sites(replace(actor, must_change_password=True))
    exported = AssetDirectory(engine).export_sites(replace(actor, permissions=frozenset({'assets.read'})))
    exported.file.close()


def test_export_uses_one_consistent_snapshot_when_site_changes_during_read(export_context):
    from sqlalchemy import event
    engine, _, actor, _ = export_context
    changed = False

    def concurrent_change(connection, cursor, statement, parameters, context, many):
        nonlocal changed
        if not changed and context.execution_options.get('stream_results'):
            changed = True
            # Boundary injection: another transaction commits after the read cursor's snapshot.
            with engine.begin() as other:
                other.execute(text("UPDATE sites SET site_name='Changed site' WHERE id=:id"), {'id': SITE_ID})
                other.execute(text("UPDATE site_product_lines SET power_mw=999 WHERE site_id=:id AND product_line='PV'"), {'id': SITE_ID})
                other.execute(text("INSERT INTO sites(id,customer_id,site_name) VALUES (:id,:customer,'Late site')"),
                              {'id': uuid4(), 'customer': CUSTOMER_ID})

    event.listen(engine, 'after_cursor_execute', concurrent_change)
    try:
        old = AssetDirectory(engine).export_sites(actor)
    finally:
        event.remove(engine, 'after_cursor_execute', concurrent_change)
    try:
        book = load_workbook(old.file, read_only=True)
        old_row = next(row for row in book['Sites'].iter_rows(max_col=13, values_only=True) if row[0] == str(SITE_ID))
        assert old_row[1] == 'Prairie Sun'
        assert old_row[6] == '100.000000'
        assert old.site_count == 2
        book.close()
    finally:
        old.file.close()
    new = AssetDirectory(engine).export_sites(actor)
    try:
        book = load_workbook(new.file, read_only=True)
        new_row = next(row for row in book['Sites'].iter_rows(max_col=13, values_only=True) if row[0] == str(SITE_ID))
        assert new_row[1] == 'Changed site'
        assert new_row[6] == '999.000000'
        assert new.site_count == 3
        book.close()
    finally:
        new.file.close()


def test_illegal_source_text_fails_safely_and_cleans_temporary_files(export_context, monkeypatch, tmp_path):
    import tempfile
    from dga.assets.public import AssetQueryError
    engine, identity, actor, _ = export_context
    monkeypatch.setattr(tempfile, 'tempdir', str(tmp_path))
    with engine.begin() as connection:
        connection.execute(text('UPDATE sites SET site_name=:name WHERE id=:id'), {'id': SITE_ID, 'name': 'Bad\x01site'})
    with pytest.raises(AssetQueryError, match='site_export_unavailable') as caught:
        AssetDirectory(engine).export_sites(actor)
    assert caught.value.status == 503
    assert not list(tmp_path.iterdir())
    assert not [event for event in identity.audit_events(actor) if event['action_code'] == 'SITE_EXPORT_PREPARED']


def test_failed_audit_never_provides_a_download_or_claims_success(export_context, monkeypatch, tmp_path):
    import tempfile
    from sqlalchemy import event
    from dga.assets.public import AssetQueryError
    engine, identity, actor, _ = export_context
    monkeypatch.setattr(tempfile, 'tempdir', str(tmp_path))

    def unavailable(connection, cursor, statement, parameters, context, many):
        if isinstance(parameters, dict) and parameters.get('action') == 'SITE_EXPORT_PREPARED':
            raise RuntimeError('Sensitive audit failure details')

    event.listen(engine, 'before_cursor_execute', unavailable)
    try:
        with pytest.raises(AssetQueryError) as caught:
            AssetDirectory(engine).export_sites(actor)
    finally:
        event.remove(engine, 'before_cursor_execute', unavailable)
    assert caught.value.code == 'site_export_unavailable'
    assert caught.value.status == 503
    assert not list(tmp_path.iterdir())
    assert not [event for event in identity.audit_events(actor) if event['action_code'] == 'SITE_EXPORT_PREPARED']
