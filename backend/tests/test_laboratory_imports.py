"""Staged laboratory originals through the approved public application seam."""
from uuid import uuid4
import pytest
from dga.assets.public import AssetDirectory
from dga.laboratory.public import LaboratoryOperations, SampleRegistry
from dga.shared.auth.public import AuditTrail
from tests.test_sample_reception import reception_context


class Files:
    def __init__(self):
        self.objects = {}

    def put(self, *, object_key, content, content_type):
        self.objects[object_key] = content

    def get(self, *, object_key):
        return self.objects[object_key]

    def delete(self, *, object_key):
        del self.objects[object_key]


def test_first_create_is_private_idempotent_and_does_not_create_formal_samples(database_url):
    from dga.laboratory.public import LaboratoryImports
    engine, identity, actor = reception_context(database_url)
    registry = SampleRegistry(engine, AssetDirectory(engine), AuditTrail())
    ops = LaboratoryOperations(engine, registry, AssetDirectory(engine), AuditTrail())
    before = ops.dashboard(actor)['metrics']
    imports = LaboratoryImports(engine, Files())
    key = str(uuid4())
    first = imports.create(actor, idempotency_key=key, format_version=1,
                           input_format='XLSX', purpose='HISTORICAL')
    again = imports.create(actor, idempotency_key=key, format_version=1,
                           input_format='XLSX', purpose='HISTORICAL')
    assert first['operation']['result'] == again['operation']['result']
    assert first['batch']['id'] == again['batch']['id']
    assert first['batch']['business_state'] == 'STAGED'
    assert first['batch']['manifest_state'] == 'EMPTY'
    assert first['batch']['revision'] == 0
    assert imports.list_batches(actor)['total'] == 1
    assert imports.operation(actor, key)['result']['batch_id'] == first['batch']['id']
    assert ops.dashboard(actor)['metrics'] == before
    assert len([e for e in identity.audit_events(actor)
                if e['action_code'] == 'LAB_IMPORT_CREATED']) == 1
    engine.dispose()



def test_source_labs_require_explicit_identity_decisions_and_retry_reuses_original_identity(database_url):
    import pytest
    from dga.laboratory.public import LaboratoryImports, LaboratoryError
    engine, _, actor = reception_context(database_url)
    imports = LaboratoryImports(engine, Files())
    empty = imports.source_labs(actor, exact_name='Source Lab')
    assert empty['exact_name_total'] == 0
    choice = dict(key='LAB_A', decision='NEW', proposed_name='Source Lab',
                  source_kind='THIRD_PARTY', reason='Confirmed source identity',
                  confirm_distinct=True, candidate_digest=empty['candidate_digest'])
    key = str(uuid4())
    first = imports.create(actor, idempotency_key=key, input_format='CSV_PAIR',
                           purpose='DAILY', source_labs=[choice])
    again = imports.create(actor, idempotency_key=key, input_format='CSV_PAIR',
                           purpose='DAILY', source_labs=[choice])
    source_id = first['batch']['source_labs'][0]['source_lab_id']
    assert again['batch']['source_labs'][0]['source_lab_id'] == source_id
    rejected_key = str(uuid4())
    with pytest.raises(LaboratoryError, match='source_candidates_changed'):
        imports.create(actor, idempotency_key=rejected_key, input_format='XLSX',
                       purpose='DAILY', source_labs=[choice])
    assert imports.operation(actor, rejected_key)['state'] == 'REJECTED'
    assert imports.list_batches(actor)['total'] == 1
    current = imports.source_labs(actor, exact_name='source lab')
    assert current['exact_name_total'] == 1
    different = imports.create(actor, idempotency_key=str(uuid4()), input_format='XLSX',
        purpose='HISTORICAL', source_labs=[choice | {'candidate_digest': current['candidate_digest']}])
    assert different['batch']['source_labs'][0]['source_lab_id'] != source_id
    referenced = imports.create(actor, idempotency_key=str(uuid4()), input_format='XLSX',
        purpose='HISTORICAL', source_labs=[dict(key='MyAlias', decision='REFERENCE',
        source_lab_id=source_id, reason='Explicit existing source')])
    assert referenced['batch']['source_labs'][0]['source_lab_id'] == source_id
    assert imports.source_labs(actor, query='SOURCE')['total'] == 2
    engine.dispose()



def test_received_original_is_confirmed_once_retained_and_private_after_reopen(database_url):
    from hashlib import sha256
    from dga.laboratory.public import LaboratoryImports
    engine, _, actor = reception_context(database_url)
    files = Files()
    imports = LaboratoryImports(engine, files)
    batch = imports.create(actor, idempotency_key=str(uuid4()), input_format='XLSX',
                           purpose='HISTORICAL')['batch']
    content = b'Original evidence retained before structured validation'
    key = str(uuid4())
    receipt = imports.receive_file(actor, batch['id'], idempotency_key=key, role='WORKBOOK',
        filename='source.xlsx', content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        content=content)
    assert receipt['operation']['result']['outcome'] == 'FILE_RECEIVED'
    assert receipt['file']['state'] == 'RECEIVED' and receipt['file']['active'] is True
    assert receipt['file']['sha256'] == sha256(content).hexdigest()
    replay = imports.receive_file(actor, batch['id'], idempotency_key=key, role='WORKBOOK',
        filename='source.xlsx', content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        content=content)
    assert replay['operation']['result'] == receipt['operation']['result']
    reopened = LaboratoryImports(engine, files).get(actor, batch['id'])
    assert reopened['revision'] == 1 and reopened['manifest_state'] == 'COMPLETE'
    assert reopened['business_state'] == 'STAGED' and reopened['validation_state'] == 'NOT_REQUESTED'
    assert imports.list_files(actor, batch['id'])['total'] == 1
    original = imports.original(actor, batch['id'], receipt['file']['id'])
    assert original['content'] == content and original['filename'] == 'source.xlsx'
    assert 'object_key' not in str(reopened) and 'lease_token' not in str(receipt)
    engine.dispose()



def test_abandoned_private_receipt_can_be_replaced_in_same_batch_by_another_current_writer(database_url):
    import pytest
    from datetime import datetime, timedelta, timezone
    from tests.identity_seed import seed_legacy_user
    from tests.test_sample_reception import INITIAL_PASSWORD, CHANGED_PASSWORD
    from dga.shared.auth.public import IdentityError
    from dga.laboratory.public import LaboratoryImports, LaboratoryError

    class Crash(BaseException):
        pass

    class CrashOnceFiles(Files):
        armed = True
        def put(self, **kwargs):
            super().put(**kwargs)
            if self.armed:
                self.armed = False
                raise Crash('Process ended after provider wrote original')

    engine, identity, admin = reception_context(database_url)
    user_id = seed_legacy_user(database_url, 'import-a', 'Uploader A', INITIAL_PASSWORD, ['lab_admin'])
    login = identity.login('import-a', INITIAL_PASSWORD)
    actor_a = identity.change_password(login.token, INITIAL_PASSWORD, CHANGED_PASSWORD).actor
    now = [datetime(2026, 10, 30, tzinfo=timezone.utc)]
    files = CrashOnceFiles()
    imports = LaboratoryImports(engine, files, clock=lambda: now[0])
    batch = imports.create(actor_a, idempotency_key=str(uuid4()), input_format='XLSX',
                           purpose='HISTORICAL')['batch']
    key_a = str(uuid4())
    values = dict(role='WORKBOOK', filename='a.xlsx',
                  content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', content=b'A evidence')
    with pytest.raises(Crash):
        imports.receive_file(actor_a, batch['id'], idempotency_key=key_a, **values)
    pending = imports.list_files(admin, batch['id'])['items'][0]
    assert pending['state'] == 'PENDING' and not pending['active']
    reject_key = str(uuid4())
    with pytest.raises(LaboratoryError, match='files_pending'):
        imports.revise(admin, batch['id'], idempotency_key=reject_key, expected_revision=0,
                       reason='Explicit selection cannot confirm pending bytes', active_files={'WORKBOOK': pending['id']})
    assert imports.operation(admin, reject_key)['state'] == 'REJECTED'
    identity.change_roles(admin, user_id, add=['field_engineer'], remove=['lab_admin'], expected_revision=0)
    with pytest.raises(IdentityError):
        imports.receive_file(actor_a, batch['id'], idempotency_key=key_a, **values)
    with pytest.raises(LaboratoryError, match='import_operation_not_found'):
        imports.operation(admin, key_a)
    own = imports.receive_file(admin, batch['id'], idempotency_key=str(uuid4()),
                              **(values | {'filename': 'b.xlsx', 'content': b'B confirmed working input'}))
    reopened = imports.get(admin, batch['id'])
    assert reopened['manifest_state'] == 'COMPLETE'
    assert [x['id'] for x in reopened['active_files']] == [own['file']['id']]
    assert b'A evidence' in files.objects.values()
    assert imports.list_batches(admin, actor_id=str(admin.user_id))['total'] == 1
    now[0] += timedelta(seconds=61)
    identity.change_roles(admin, user_id, add=['lab_admin'], remove=['field_engineer'], expected_revision=1)
    late = imports.receive_file(actor_a, batch['id'], idempotency_key=key_a, **values)
    assert late['file']['id'] == pending['id'] and not late['file']['active']
    current = imports.get(admin, batch['id'])
    assert current['revision'] == 2 and current['manifest_state'] == 'COMPLETE'
    assert current['active_files'][0]['id'] == own['file']['id']
    assert imports.original(admin, batch['id'], own['file']['id'])['content'] == b'B confirmed working input'
    assert imports.revision_history(admin, batch['id'])['total'] == 2
    engine.dispose()



def test_versioned_templates_have_complete_two_table_schema_and_honest_synthetic_examples(database_url):
    from io import BytesIO, StringIO
    import csv
    import json
    from zipfile import ZipFile
    from openpyxl import load_workbook
    from dga.laboratory.public import LaboratoryImports
    engine, _, actor = reception_context(database_url)
    imports = LaboratoryImports(engine, Files())
    expected_samples = 'source_lab_key,source_lab_name,external_sample_id,source_site_name,source_whole_unit_serial,source_transformer_serial,formal_asset_id,system_asset_number,sampled_at,sampled_timezone,sampling_context,physical_received,received_at,received_timezone,container_count,report_filenames,notes,source_note'.split(',')
    expected_tests = 'source_lab_key,external_sample_id,external_test_id,test_type,measured_at,measured_timezone,source_method,source_method_version,method_version_id,report_selected,source_instrument,source_analyst,source_qa,h2_value,h2_qualifier,h2_unit,ch4_value,ch4_qualifier,ch4_unit,c2h2_value,c2h2_qualifier,c2h2_unit,c2h4_value,c2h4_qualifier,c2h4_unit,c2h6_value,c2h6_qualifier,c2h6_unit,co_value,co_qualifier,co_unit,co2_value,co2_qualifier,co2_unit,result_value,result_qualifier,result_unit,notes,source_note'.split(',')
    xlsx = imports.template(actor, format='XLSX')
    book = load_workbook(BytesIO(xlsx['content']))
    assert set(book.sheetnames) == {'Format', 'Instructions', '油样信息', '检测结果'}
    assert list(book['Format'].values) == [('version', 'format'), (1, 'LABORATORY')]
    assert list(next(book['油样信息'].values)) == expected_samples
    assert list(next(book['检测结果'].values)) == expected_tests
    assert book['油样信息'].max_row == book['检测结果'].max_row == 1
    example = imports.template(actor, format='CSV_PAIR', example=True)
    with ZipFile(BytesIO(example['content'])) as package:
        assert set(package.namelist()) == {'samples.csv', 'tests.csv', 'instructions.json'}
        samples = list(csv.DictReader(StringIO(package.read('samples.csv').decode('utf-8'))))
        tests = list(csv.DictReader(StringIO(package.read('tests.csv').decode('utf-8'))))
        info = json.loads(package.read('instructions.json'))
    assert list(samples[0]) == expected_samples and list(tests[0]) == expected_tests
    assert samples[0]['external_sample_id'] == '00017'
    assert len(tests) == 2 and tests[0]['external_test_id'] != tests[1]['external_test_id']
    assert tests[0]['method_version_id'] == '' and samples[0]['formal_asset_id'] == ''
    assert info['version'] == 1 and info['example_is_synthetic'] is True
    assert imports.list_batches(actor)['total'] == 0
    engine.dispose()



def test_rejected_create_does_not_expose_private_operation_key_in_shared_audit(database_url):
    import pytest
    from dga.laboratory.public import LaboratoryImports, LaboratoryError
    engine, identity, actor = reception_context(database_url)
    imports = LaboratoryImports(engine, Files())
    key = str(uuid4())
    with pytest.raises(LaboratoryError, match='source_candidates_changed'):
        imports.create(actor, idempotency_key=key, input_format='XLSX', purpose='DAILY',
            source_labs=[dict(key='SRC', decision='NEW', proposed_name='New source',
                source_kind='THIRD_PARTY', reason='Distinct identity', confirm_distinct=True,
                candidate_digest='0'*64)])
    assert imports.operation(actor, key)['state'] == 'REJECTED'
    assert key not in str(identity.audit_events(actor))
    engine.dispose()


def test_unknown_create_and_receipt_commits_recover_without_false_receipts_or_deletion(database_url):
    from contextlib import contextmanager
    from datetime import datetime, timedelta, timezone
    import pytest
    from dga.laboratory.public import LaboratoryImports, LaboratoryError

    class FaultEngine:
        def __init__(self, engine):
            self.engine, self.fault = engine, None
        def arm(self, number, committed):
            self.calls, self.fault = 0, (number, committed)
        @contextmanager
        def begin(self):
            self.calls = getattr(self, 'calls', 0) + 1
            fault = self.fault if self.fault and self.calls == self.fault[0] else None
            with self.engine.begin() as c:
                yield c
                if fault and not fault[1]:
                    self.fault = None
                    raise ConnectionError('COMMIT not confirmed; database rolled back')
            if fault and fault[1]:
                self.fault = None
                raise ConnectionError('COMMIT response lost after durable database commit')

    engine, _, actor = reception_context(database_url)
    boundary, files = FaultEngine(engine), Files()
    now = [datetime(2026, 10, 30, tzinfo=timezone.utc)]
    imports = LaboratoryImports(boundary, files, clock=lambda: now[0])
    for committed in (False, True):
        key = str(uuid4())
        boundary.arm(1, committed)
        with pytest.raises(ConnectionError):
            imports.create(actor, idempotency_key=key, input_format='XLSX', purpose='DAILY')
        if committed:
            assert imports.operation(actor, key)['state'] == 'COMPLETED'
        else:
            with pytest.raises(LaboratoryError, match='import_operation_not_found'):
                imports.operation(actor, key)
        batch = imports.create(actor, idempotency_key=key, input_format='XLSX', purpose='DAILY')['batch']
        for commit_number in (1, 2):  # reservation then final receipt transaction
            for durable in (False, True):
                upload_key = str(uuid4())
                content = f'{committed}/{commit_number}/{durable}'.encode()
                values = dict(role='WORKBOOK', filename='source.xlsx',
                    content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', content=content)
                before = dict(files.objects)
                boundary.arm(commit_number, durable)
                with pytest.raises(ConnectionError):
                    imports.receive_file(actor, batch['id'], idempotency_key=upload_key, **values)
                if commit_number == 1:
                    assert files.objects == before  # no write before reservation COMMIT is known
                else:
                    assert content in files.objects.values()  # never delete an uncertain final receipt
                if commit_number == 1 and not durable:
                    with pytest.raises(LaboratoryError, match='import_operation_not_found'):
                        imports.operation(actor, upload_key)
                else:
                    operation = imports.operation(actor, upload_key)
                    assert operation['state'] == ('COMPLETED' if commit_number == 2 and durable else 'PENDING')
                    if operation['state'] == 'PENDING':
                        assert operation['result'] is None
                now[0] += timedelta(seconds=61)
                receipt = imports.receive_file(actor, batch['id'], idempotency_key=upload_key, **values)
                assert receipt['operation']['state'] == 'COMPLETED'
                assert imports.original(actor, batch['id'], receipt['file']['id'])['content'] == content
    assert imports.list_batches(actor)['total'] == 2
    engine.dispose()



def test_history_filters_pagination_and_revision_replay_use_staged_evidence(database_url):
    from dga.laboratory.public import LaboratoryImports, LaboratoryError
    import pytest
    engine, _, actor = reception_context(database_url)
    imports = LaboratoryImports(engine, Files())
    batches = [imports.create(actor, idempotency_key=str(uuid4()), input_format='CSV_PAIR',
               purpose='DAILY' if i == 20 else 'HISTORICAL')['batch'] for i in range(21)]
    assert imports.list_batches(actor, purpose='DAILY')['total'] == 1
    assert imports.list_batches(actor, q=batches[0]['id'])['total'] == 1
    assert len(imports.list_batches(actor, page=1)['items']) == 20
    assert len(imports.list_batches(actor, page=2)['items']) == 1
    key = str(uuid4())
    result = imports.revise(actor, batches[0]['id'], idempotency_key=key,
        expected_revision=0, reason='Confirm source context',
        metadata=dict(source_timezone='America/Chicago', source_timezone_confirmed=True))
    assert result['batch']['revision'] == 1
    assert imports.revise(actor, batches[0]['id'], idempotency_key=key, expected_revision=0,
        reason='Confirm source context', metadata=dict(source_timezone='America/Chicago',
        source_timezone_confirmed=True))['operation'] == result['operation']
    assert imports.revision_history(actor, batches[0]['id'])['total'] == 1
    assert imports.list_batches(actor, actor_id=str(actor.user_id))['total'] == 21
    assert imports.list_batches(actor, created_from='2099-01-01')['total'] == 0
    with pytest.raises(LaboratoryError, match='invalid_import_input'):
        imports.list_batches(actor, mystery='ignore')
    engine.dispose()


def test_parallel_exact_intents_create_one_identity_and_receive_once(database_url):
    from concurrent.futures import ThreadPoolExecutor
    from dga.laboratory.public import LaboratoryImports, LaboratoryError
    import pytest
    engine, identity, actor = reception_context(database_url)
    files = Files()
    imports = LaboratoryImports(engine, files)
    key = str(uuid4())
    choices = [dict(key='Source', decision='NEW', proposed_name='Laboratory Ω',
        source_kind='IN_HOUSE', reason='Explicit distinct source', confirm_distinct=True,
        candidate_digest=imports.source_labs(actor, exact_name='Laboratory Ω')['candidate_digest'])]
    def create(_):
        return imports.create(actor, idempotency_key=key, input_format='XLSX', purpose='DAILY',
            source_labs=choices)
    with ThreadPoolExecutor(max_workers=4) as pool:
        results = list(pool.map(create, range(4)))
    assert len({x['batch']['id'] for x in results}) == 1
    assert imports.source_labs(actor)['total'] == imports.list_batches(actor)['total'] == 1
    batch = results[0]['batch']
    receipt_key = str(uuid4())
    values = dict(role='WORKBOOK', filename='原始 Ω.xlsx',
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', content=b'evidence')
    def receive(_):
        return imports.receive_file(actor, batch['id'], idempotency_key=receipt_key, **values)
    with ThreadPoolExecutor(max_workers=4) as pool:
        receipts = list(pool.map(receive, range(4)))
    assert imports.get(actor, batch['id'])['revision'] == 1
    assert imports.list_files(actor, batch['id'])['total'] == 1
    assert len(files.objects) == 1
    assert len([e for e in identity.audit_events(actor) if e['action_code']=='LAB_IMPORT_FILE_RECEIVED']) == 1
    with pytest.raises(LaboratoryError, match='import_operation_conflict'):
        imports.receive_file(actor, batch['id'], idempotency_key=receipt_key,
            **(values | {'content':b'changed'}))
    engine.dispose()



def test_storage_failures_keep_retry_metadata_and_do_not_audit_corrupt_download_success(database_url):
    import pytest
    from dga.shared.files import ObjectStorageError
    from dga.laboratory.public import LaboratoryImports
    class FaultFiles(Files):
        fail = True
        def put(self, **values):
            if self.fail:
                raise OSError('secret provider connection information')
            super().put(**values)
    engine, identity, actor = reception_context(database_url)
    files = FaultFiles()
    imports = LaboratoryImports(engine, files)
    batch = imports.create(actor,idempotency_key=str(uuid4()),input_format='CSV_PAIR',purpose='DAILY')['batch']
    key = str(uuid4())
    values = dict(role='SAMPLES',filename='samples.csv',content_type='text/csv',content=b'raw,bytes\r\n')
    with pytest.raises(ObjectStorageError):
        imports.receive_file(actor,batch['id'],idempotency_key=key,**values)
    assert imports.operation(actor,key)['state']=='RETRYABLE'
    assert imports.get(actor,batch['id'])['revision']==0
    pending = imports.list_files(actor,batch['id'])['items'][0]
    assert pending['state']=='RETRYABLE' and not pending['active']
    files.fail=False
    received=imports.receive_file(actor,batch['id'],idempotency_key=key,**values)
    assert received['file']['id']==pending['id']
    assert imports.get(actor,batch['id'])['manifest_state']=='PARTIAL'
    before=len([e for e in identity.audit_events(actor) if e['action_code']=='LAB_IMPORT_SOURCE_DOWNLOADED'])
    files.objects[next(iter(files.objects))]=b'wrong'
    with pytest.raises(ObjectStorageError):
        imports.original(actor,batch['id'],pending['id'])
    assert len([e for e in identity.audit_events(actor) if e['action_code']=='LAB_IMPORT_SOURCE_DOWNLOADED'])==before
    engine.dispose()


def test_current_authority_covers_every_staging_read_and_write_without_changing_analyst_policy(database_url):
    import pytest
    from tests.identity_seed import seed_legacy_user
    from tests.test_sample_reception import INITIAL_PASSWORD, CHANGED_PASSWORD
    from dga.shared.auth.public import IdentityError
    from dga.laboratory.public import LaboratoryImports
    engine, identity, admin = reception_context(database_url)
    imports = LaboratoryImports(engine, Files())
    user_id=seed_legacy_user(database_url,'receipt-analyst','Analyst',INITIAL_PASSWORD,['analyst'])
    login=identity.login('receipt-analyst',INITIAL_PASSWORD)
    actor=identity.change_password(login.token,INITIAL_PASSWORD,CHANGED_PASSWORD).actor
    assert 'laboratory.finalize' in actor.permissions
    key=str(uuid4())
    batch=imports.create(actor,idempotency_key=key,input_format='XLSX',purpose='DAILY')['batch']
    upload=dict(role='WORKBOOK',filename='data.xlsx',
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',content=b'evidence')
    file=imports.receive_file(actor,batch['id'],idempotency_key=str(uuid4()),**upload)['file']
    identity.change_roles(admin,user_id,add=['field_engineer'],remove=['analyst'],expected_revision=0)
    calls=[
        lambda: imports.template(actor),lambda: imports.source_labs(actor),
        lambda: imports.reject_input(actor,idempotency_key=str(uuid4()),action='CREATE',request=dict(input_format='XLSX',purpose='DAILY',source_timezone='Invalid/Zone')),
        lambda: imports.list_batches(actor),lambda: imports.get(actor,batch['id']),
        lambda: imports.list_files(actor,batch['id']),lambda: imports.operation(actor,key),
        lambda: imports.original(actor,batch['id'],file['id']),
        lambda: imports.revision_history(actor,batch['id']),
        lambda: imports.create(actor,idempotency_key=str(uuid4()),input_format='XLSX',purpose='DAILY'),
        lambda: imports.receive_file(actor,batch['id'],idempotency_key=str(uuid4()),**upload),
        lambda: imports.revise(actor,batch['id'],idempotency_key=str(uuid4()),expected_revision=1,
            reason='Stale actor denied',metadata=dict(purpose='HISTORICAL'))]
    for call in calls:
        with pytest.raises(IdentityError):
            call()
    reader=identity.login('receipt-analyst',CHANGED_PASSWORD).actor
    assert {'assets.read','laboratory.read','analysis.read'} <= reader.permissions
    with pytest.raises(IdentityError):
        imports.get(reader,batch['id'])
    assert imports.get(admin,batch['id'])['revision']==1
    engine.dispose()


def test_receipt_input_limits_are_enforced_before_provider_or_reservation(database_url):
    import pytest
    from dga.laboratory.public import LaboratoryImports, LaboratoryError
    engine, _, actor=reception_context(database_url)
    files=Files()
    imports=LaboratoryImports(engine,files)
    batch=imports.create(actor,idempotency_key=str(uuid4()),input_format='XLSX',purpose='DAILY')['batch']
    base=dict(role='WORKBOOK',filename='source.xlsx',
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',content=b'bytes')
    for patch in ({'filename':'../source.xlsx'},{'filename':'C:source.xlsx'},{'filename':'bad\n.xlsx'},
        {'filename':'a'*256+'.xlsx'},{'content_type':'text/plain'},{'role':'SAMPLES','filename':'a.csv','content_type':'text/csv'},
        {'content':b''},{'content':b'a'*(10*1024*1024+1)}):
        with pytest.raises(LaboratoryError):
            imports.receive_file(actor,batch['id'],idempotency_key=str(uuid4()),**(base|patch))
    assert not files.objects and imports.list_files(actor,batch['id'])['total']==0
    exact=imports.receive_file(actor,batch['id'],idempotency_key=str(uuid4()),**(base|{'filename':'原'*250+'.xlsx','content':b'x'*(10*1024*1024)}))
    assert exact['file']['byte_size']==10*1024*1024 and len(exact['file']['filename'])==255
    engine.dispose()


def test_expired_upload_token_cannot_complete_after_new_claim_receives_same_original(database_url):
    import pytest
    from datetime import datetime, timedelta, timezone
    from concurrent.futures import ThreadPoolExecutor
    from threading import Event
    from dga.laboratory.public import LaboratoryImports, LaboratoryError
    entered, release=Event(), Event()
    class PausingFiles(Files):
        first=True
        def put(self,**values):
            super().put(**values)
            if self.first:
                self.first=False
                entered.set()
                assert release.wait(15)
    engine, _, actor=reception_context(database_url)
    files=PausingFiles()
    now=[datetime(2026,10,30,tzinfo=timezone.utc)]
    imports=LaboratoryImports(engine,files,clock=lambda:now[0])
    batch=imports.create(actor,idempotency_key=str(uuid4()),input_format='XLSX',purpose='DAILY')['batch']
    key=str(uuid4())
    values=dict(role='WORKBOOK',filename='source.xlsx',
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',content=b'exact evidence')
    with ThreadPoolExecutor(max_workers=1) as pool:
        first=pool.submit(imports.receive_file,actor,batch['id'],idempotency_key=key,**values)
        try:
            assert entered.wait(10)
            now[0]+=timedelta(seconds=61)
            recovered=imports.receive_file(actor,batch['id'],idempotency_key=key,**values)
        finally:
            release.set()
        with pytest.raises(LaboratoryError,match='import_receipt_claim_stale'):
            first.result()
    assert recovered['operation']['result']['applied_revision']==1
    assert imports.get(actor,batch['id'])['revision']==1
    assert imports.list_files(actor,batch['id'])['total']==1
    engine.dispose()



def test_failed_receipt_audit_rolls_back_completion_retains_original_and_can_recover(database_url):
    from datetime import datetime, timedelta, timezone
    import pytest
    from dga.laboratory.public import LaboratoryImports
    from dga.shared.auth.public import AuditTrail
    class FaultAudit(AuditTrail):
        fail=True
        def append(self,connection,actor,action,**values):
            if self.fail and str(action)=='LAB_IMPORT_FILE_RECEIVED':
                raise ConnectionError('Audit response failed inside final transaction')
            return super().append(connection,actor,action,**values)
    engine,identity,actor=reception_context(database_url)
    files=Files()
    now=[datetime(2026,10,30,tzinfo=timezone.utc)]
    audit=FaultAudit(clock=lambda:now[0])
    imports=LaboratoryImports(engine,files,audit=audit,clock=lambda:now[0])
    batch=imports.create(actor,idempotency_key=str(uuid4()),input_format='XLSX',purpose='DAILY')['batch']
    key=str(uuid4())
    values=dict(role='WORKBOOK',filename='original.xlsx',
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',content=b'original retained')
    with pytest.raises(ConnectionError):
        imports.receive_file(actor,batch['id'],idempotency_key=key,**values)
    assert b'original retained' in files.objects.values()
    assert imports.operation(actor,key)['state']=='PENDING'
    assert imports.get(actor,batch['id'])['revision']==0
    assert imports.revision_history(actor,batch['id'])['total']==0
    assert not [e for e in identity.audit_events(actor) if e['action_code']=='LAB_IMPORT_FILE_RECEIVED']
    audit.fail=False
    now[0]+=timedelta(seconds=61)
    done=imports.receive_file(actor,batch['id'],idempotency_key=key,**values)
    assert done['operation']['result']['applied_revision']==1
    assert imports.list_files(actor,batch['id'])['total']==1
    engine.dispose()




def test_history_summaries_describe_active_manifest_without_loading_business_rows(database_url):
    from dga.laboratory.public import LaboratoryImports
    engine,_,actor=reception_context(database_url)
    imports=LaboratoryImports(engine,Files())
    batch=imports.create(actor,idempotency_key=str(uuid4()),input_format='CSV_PAIR',purpose='DAILY')['batch']
    assert imports.list_batches(actor)['items'][0]['manifest_state']=='EMPTY'
    for role in ('SAMPLES','TESTS'):
        imports.receive_file(actor,batch['id'],idempotency_key=str(uuid4()),role=role,
            filename=role.lower()+'.csv',content_type='text/csv',content=b'raw evidence')
    summary=imports.list_batches(actor)['items'][0]
    assert summary['manifest_state']=='COMPLETE'
    assert summary['received_file_count']==summary['active_file_count']==2
    assert summary['validation_state']=='NOT_REQUESTED' and summary['business_state']=='STAGED'
    assert 'rows' not in summary and 'object_key' not in str(summary)
    engine.dispose()




@pytest.mark.parametrize('database_url',['head'],indirect=True)
def test_unicode_filename_history_search_uses_one_comparison_and_complete_pages(database_url):
    """PostgreSQL lower on both sides; sharp-s and ss stay distinct filenames."""
    from dga.laboratory.public import LaboratoryImports
    engine,_,actor=reception_context(database_url)
    imports=LaboratoryImports(engine,Files())
    expected=[]
    def original(filename):
        batch=imports.create(actor,idempotency_key=str(uuid4()),input_format='CSV_PAIR',purpose='DAILY')['batch']
        imports.receive_file(actor,batch['id'],idempotency_key=str(uuid4()),role='SAMPLES',
            filename=filename,content_type='text/csv',content=b'retained evidence')
        return batch['id']
    for _ in range(21):
        expected.append(original('Straße.csv'))
    ss=original('STRASSE.csv')
    chinese=original('来源油样.csv')
    ascii_id=original('Mixed-CASE.csv')
    first=imports.list_batches(actor,q='Straße.csv')
    second=imports.list_batches(actor,q='STRAße.CSV',page=2)
    assert first['total']==second['total']==21
    assert len(first['items'])==20 and len(second['items'])==1
    assert {x['id'] for x in first['items']+second['items']}==set(expected)
    assert imports.list_batches(actor,q='strasse.csv')['total']==1
    assert imports.list_batches(actor,q='strasse.csv')['items'][0]['id']==ss
    assert imports.list_batches(actor,q='来源油样.csv')['items'][0]['id']==chinese
    assert imports.list_batches(actor,q='mixed-case.CSV')['items'][0]['id']==ascii_id
    assert imports.list_files(actor,expected[0])['items'][0]['filename']=='Straße.csv'
    assert imports.list_batches(actor,q=expected[0].upper())['items'][0]['id']==expected[0]
    engine.dispose()



@pytest.mark.parametrize('database_url',['head'],indirect=True)
def test_rejected_input_proof_requires_commit_and_audit_and_preserves_private_intent(database_url):
    from contextlib import contextmanager
    from dga.laboratory.public import LaboratoryImports, LaboratoryError
    from dga.shared.auth.public import AuditTrail

    class FaultEngine:
        def __init__(self, engine):
            self.engine, self.armed = engine, False
        def arm(self, committed):
            self.armed, self.committed = True, committed
        @contextmanager
        def begin(self):
            fault, committed = self.armed, getattr(self, 'committed', False)
            self.armed = False
            with self.engine.begin() as c:
                yield c
                if fault and not committed:
                    raise ConnectionError('Rejected-input transaction rolled back')
            if fault and committed:
                raise ConnectionError('Rejected-input COMMIT response lost')

    engine,identity,actor=reception_context(database_url)
    boundary=FaultEngine(engine)
    imports=LaboratoryImports(boundary,Files())
    request=dict(input_format='XLSX',purpose='DAILY',source_timezone='america/chicago')
    for committed in (False,True):
        key=str(uuid4())
        boundary.arm(committed)
        with pytest.raises(ConnectionError):
            imports.reject_input(actor,idempotency_key=key,action='CREATE',request=request)
        if committed:
            assert imports.operation(actor,key)['state']=='REJECTED'
        else:
            with pytest.raises(LaboratoryError,match='import_operation_not_found'):
                imports.operation(actor,key)
        proof=imports.reject_input(actor,idempotency_key=key,action='CREATE',request=request)
        assert proof['state']=='REJECTED'
        assert imports.reject_input(actor,idempotency_key=key,action='CREATE',request=request)==proof
        assert key not in str(identity.audit_events(actor))
        with pytest.raises(LaboratoryError,match='import_operation_conflict'):
            imports.reject_input(actor,idempotency_key=key,action='CREATE',request=request|{'other':'changed'})
    with pytest.raises(LaboratoryError,match='import_operation_conflict'):
        imports.reject_input(actor,idempotency_key=key,action='REVISE',batch_id=str(uuid4()),
            request=dict(expected_revision=True,reason='Invalid',metadata=dict(purpose='DAILY')))
    assert len([e for e in identity.audit_events(actor) if e['action_code']=='LAB_IMPORT_CREATED'])==2
    class FaultAudit(AuditTrail):
        def append(self,*args,**kwargs):
            raise ConnectionError('Audit cannot commit')
    audit_failure=LaboratoryImports(engine,Files(),audit=FaultAudit())
    audit_key=str(uuid4())
    with pytest.raises(ConnectionError):
        audit_failure.reject_input(actor,idempotency_key=audit_key,action='CREATE',request=request)
    with pytest.raises(LaboratoryError,match='import_operation_not_found'):
        imports.operation(actor,audit_key)
    valid_key=str(uuid4())
    with pytest.raises(LaboratoryError,match='invalid_import_input'):
        imports.reject_input(actor,idempotency_key=valid_key,action='CREATE',
            request=dict(input_format='XLSX',purpose='DAILY'))
    with pytest.raises(LaboratoryError,match='import_operation_not_found'):
        imports.operation(actor,valid_key)
    assert imports.list_batches(actor)['total']==0
    engine.dispose()



