"""Upgrade both actual deployed main baselines without losing public behavior."""
from io import BytesIO
from uuid import uuid4

import pytest
from alembic import command
from alembic.config import Config
from openpyxl import load_workbook
from sqlalchemy import create_engine, text

from dga.assets.public import AssetDirectory
from dga.shared.auth.public import IdentityService
from tests.identity_seed import seed_legacy_user

INITIAL = 'Sync initial isolated passphrase 63!'
CHANGED = 'Sync changed isolated passphrase 63!'


@pytest.mark.parametrize('database_url', ['0025_site_export_audit', '0029_oidc_callback'], indirect=True)
def test_both_main_baselines_keep_accounts_sites_and_distinct_audit_payloads(database_url, request):
    baseline = request.node.callspec.params['database_url']
    engine = create_engine(database_url)
    customer, site, event_id = uuid4(), uuid4(), uuid4()
    try:
        user_id = seed_legacy_user(database_url, 'sync-reader', 'Retained administrator', INITIAL, ['system_admin'])
        identity = IdentityService(engine)
        first = identity.login('sync-reader', INITIAL)
        identity.change_password(first.token, INITIAL, CHANGED)
        with engine.begin() as connection:
            connection.execute(text("INSERT INTO customers(id,customer_name) VALUES (:id,'Retained customer')"), {'id': customer})
            connection.execute(text("INSERT INTO sites(id,customer_id,site_name,location_text) VALUES (:id,:customer,'Retained site','Texas')"), {'id': site, 'customer': customer})
            connection.execute(text("INSERT INTO site_product_lines(site_id,product_line,power_mw) VALUES (:id,'PV',1234567890.123456)"), {'id': site})
            if baseline == '0025_site_export_audit':
                connection.execute(text("""INSERT INTO audit_logs(id,actor_user_id,action_code,result,occurred_at,details)
                    VALUES (:id,:user,'RETAINED_EXPORT','SUCCESS',now(),CAST(:details AS jsonb))"""), {'id': event_id, 'user': user_id, 'details': '{"scope":"retained","site_count":1}'})
            else:
                connection.execute(text("""INSERT INTO audit_logs(id,actor_user_id,action_code,result,occurred_at,before_value,after_value)
                    VALUES (:id,:user,'RETAINED_IDENTITY','SUCCESS',now(),'{"status":"before"}','{"status":"after"}')"""), {'id': event_id, 'user': user_id})
        command.upgrade(Config('alembic.ini'), 'head')
        actor = identity.login('sync-reader', CHANGED).actor
        assert actor.user_id == user_id and 'system_admin' in actor.roles
        exported = AssetDirectory(engine).export_sites(actor)
        try:
            book = load_workbook(exported.file, read_only=True)
            rows = list(book['Sites'].iter_rows(max_col=13, values_only=True))
            assert len(rows) == 2
            assert rows[1][0:5] == (str(site), 'Retained site', str(customer), 'Retained customer', 'Texas')
            assert rows[1][6] == '1234567890.123456'
            book.close()
        finally:
            exported.file.close()
        old = next(event for event in identity.audit_events(actor) if event['id'] == event_id)
        if baseline == '0025_site_export_audit':
            assert old['details'] == {'scope': 'retained', 'site_count': 1}
            assert old['before_value'] is None and old['after_value'] is None
        else:
            assert old['details'] == {}
            assert old['before_value'] == {'status': 'before'} and old['after_value'] == {'status': 'after'}
    finally:
        engine.dispose()
