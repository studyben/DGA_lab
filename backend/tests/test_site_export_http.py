from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import load_workbook
import pytest

from dga.main import create_app
from dga.shared.config import Settings
from tests.test_site_export import export_context
from tests.test_sample_reception import CHANGED_PASSWORD

ORIGIN = {'Origin': 'http://127.0.0.1:8080'}


@pytest.fixture
def export_client(export_context, database_url):
    with TestClient(create_app(Settings(database_url=database_url, cookie_secure=False))) as client:
        login = client.post('/api/auth/login', json={
            'username': 'reception-admin', 'password': CHANGED_PASSWORD}, headers=ORIGIN)
        assert login.status_code == 200
        yield client


def test_authenticated_download_is_a_complete_unfiltered_xlsx(export_client):
    response = export_client.get('/api/assets/sites/export?product_line=PV&page=2&site=missing')
    assert response.status_code == 200
    assert response.headers['content-type'] == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    assert response.headers['content-disposition'].startswith('attachment; filename="site-reference-')
    assert response.headers['cache-control'] == 'no-store'
    assert response.headers['x-content-type-options'] == 'nosniff'
    with BytesIO(response.content) as source:
        book = load_workbook(source, read_only=True)
        rows = list(book['Sites'].iter_rows(max_col=13, values_only=True))
        assert len(rows) == 3
        assert {row[1] for row in rows[1:]} == {'Prairie Sun', 'Empty site'}
        assert dict(book['Instructions'].values)['site_count'] == 2
        book.close()


def test_anonymous_and_read_only_permission_contract(export_context, database_url):
    engine, identity, admin, _ = export_context
    from tests.test_sample_reception import INITIAL_PASSWORD
    identity.provision_user(admin, 'export-reader', 'Reader', INITIAL_PASSWORD, ['field_engineer'])
    first = identity.login('export-reader', INITIAL_PASSWORD)
    identity.change_password(first.token, INITIAL_PASSWORD, CHANGED_PASSWORD)
    with TestClient(create_app(Settings(database_url=database_url, cookie_secure=False))) as client:
        assert client.get('/api/assets/sites/export').status_code == 401
        result = client.post('/api/auth/login', json={'username': 'export-reader', 'password': CHANGED_PASSWORD}, headers=ORIGIN)
        assert result.status_code == 200
        assert 'assets.import' not in result.json()['permissions']
        assert client.get('/api/assets/sites/export').status_code == 200
        # Arrange a separate role with no read grants; verify authorization through HTTP.
        from sqlalchemy import text
        with engine.begin() as connection:
            connection.execute(text("INSERT INTO roles(id,role_code,role_name) VALUES (gen_random_uuid(),'export_denied','No asset access')"))
        try:
            identity.set_roles(admin, first.actor.user_id, ['export_denied'])
            assert client.get('/api/assets/sites/export').status_code == 403
        finally:
            identity.set_roles(admin, first.actor.user_id, ['field_engineer'])
            with engine.begin() as connection:
                connection.execute(text("DELETE FROM roles WHERE role_code='export_denied'"))


def test_workbook_failure_is_safe_json_not_a_partial_download(export_client, export_context):
    from sqlalchemy import text
    engine, identity, actor, _ = export_context
    with engine.begin() as connection:
        connection.execute(text('UPDATE sites SET site_name=:name'), {'name': 'bad\x01name'})
    response = export_client.get('/api/assets/sites/export')
    assert response.status_code == 503
    assert response.json() == {'code': 'site_export_unavailable'}
    assert 'content-disposition' not in response.headers
    assert not [e for e in identity.audit_events(actor) if e['action_code'] == 'SITE_EXPORT_PREPARED']


@pytest.mark.parametrize('disconnect', [False, True])
def test_http_response_closes_temporary_download_even_on_disconnect(export_client, monkeypatch, disconnect):
    import asyncio
    import tempfile
    from starlette.requests import ClientDisconnect
    closed = []
    original = tempfile.SpooledTemporaryFile.close

    def close(resource):
        if not resource.closed:
            closed.append(resource)
        original(resource)

    monkeypatch.setattr(tempfile.SpooledTemporaryFile, 'close', close)
    if not disconnect:
        assert export_client.get('/api/assets/sites/export').status_code == 200
    else:
        scope = {'type': 'http', 'asgi': {'version': '3.0', 'spec_version': '2.4'},
                 'http_version': '1.1', 'method': 'GET', 'scheme': 'http',
                 'path': '/api/assets/sites/export', 'raw_path': b'/api/assets/sites/export',
                 'query_string': b'', 'root_path': '', 'server': ('testserver', 80),
                 'client': ('testclient', 50000),
                 'headers': [(b'cookie', f"dga_session={export_client.cookies.get('dga_session')}".encode())]}

        async def receive():
            return {'type': 'http.request', 'body': b'', 'more_body': False}

        async def send(message):
            if message['type'] == 'http.response.body':
                raise OSError('Client disconnected')

        with pytest.raises(ClientDisconnect):
            asyncio.run(export_client.app(scope, receive, send))
    assert len(closed) == 1
    assert closed[0].closed

