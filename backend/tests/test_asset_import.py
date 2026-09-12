from datetime import datetime, timezone
from io import BytesIO
from zipfile import ZipFile, ZIP_DEFLATED

from openpyxl import Workbook
from sqlalchemy import text
import pytest
from dataclasses import replace
from dga.shared.auth.public import IdentityError
from dga.shared.files import UnavailableFileStore, ObjectStorageError
from dga.assets.public import AssetQueryError

from dga.assets.public import AssetImports, AssetLifecycle
from tests.test_sample_reception import reception_context, CUSTOMER_ID, SITE_ID


class MemoryFiles:
    def __init__(self):
        self.files = {}

    def put(self, *, object_key, content, content_type):
        self.files[object_key] = content

    def delete(self, *, object_key):
        self.files.pop(object_key, None)


def workbook(rows):
    book = Workbook()
    sheet = book.active
    sheet.title = 'Assets'
    columns = ['record_key', 'serial_number', 'material_number', 'product_line',
               'machine_type', 'status', 'location_kind', 'effective_at',
               'customer_id', 'site_id', 'parent_record_key']
    for row in rows:
        columns.extend(k for k in row if k not in columns)
    sheet.append(columns)
    for row in rows:
        sheet.append([row.get(k) for k in columns])
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def new_rows():
    return [dict(record_key='root', serial_number='00001', material_number='IMPORT-UNIT',
                 product_line='PV', machine_type='INVERTER_UNIT', status='IN_SERVICE',
                 location_kind='SITE', effective_at='2025-01-01T08:00:00Z',
                 customer_id=str(CUSTOMER_ID), site_id=str(SITE_ID)),
            dict(record_key='tx', serial_number='00002', material_number='IMPORT-TX',
                 product_line='PV', machine_type='TRANSFORMER', status='IN_SERVICE',
                 location_kind='PARENT', effective_at='2025-01-01T09:00:00Z', parent_record_key='root')]


def import_context(database_url):
    engine, _, actor = reception_context(database_url)
    # Arrangement of known master data, not behavioral assertions against private tables.
    with engine.begin() as c:
        c.execute(text("DELETE FROM asset_materials WHERE material_number IN ('IMPORT-UNIT','IMPORT-TX')"))
        c.execute(text("INSERT INTO asset_materials VALUES ('IMPORT-UNIT','SG4400UD','WHOLE_UNIT'),('IMPORT-TX','TX-4400','TRANSFORMER')"))
        c.execute(text("INSERT INTO site_product_lines(site_id,product_line) VALUES (:site,'PV')"), {'site': SITE_ID})
    files = MemoryFiles()
    imports = AssetImports(engine, files, clock=lambda: datetime(2026, 1, 1, tzinfo=timezone.utc))
    return engine, actor, imports, files


def test_uploaded_hierarchy_is_previewed_without_publishing_assets(database_url):
    engine, actor, imports, files = import_context(database_url)
    before = AssetLifecycle(engine).catalog(actor)['total']
    content = workbook(new_rows())
    batch = imports.submit(actor, filename='首批.xlsx', content=content)
    assert batch['state'] == 'STAGED'
    assert content in files.files.values()
    assert imports.validate_next() is True
    preview = imports.get(actor, batch['id'])
    assert preview['state'] == 'VALIDATED'
    assert preview['summary'] == {'total_rows': 2, 'valid_rows': 2, 'warning_rows': 0, 'error_rows': 0, 'warning_count': 0, 'error_count': 0}
    assert preview['rows'][0]['values']['serial_number'] == '00001'
    assert AssetLifecycle(engine).catalog(actor)['total'] == before
    engine.dispose()


@pytest.mark.parametrize('change,field', [
    ({'material_number': 'MISSING'}, 'material_number'),
    ({'serial_number': 123}, 'serial_number'),
    ({'effective_at': '2025-01-01'}, 'effective_at'),
    ({'effective_at': '2027-01-01T00:00:00Z'}, 'effective_at'),
    ({'customer_id': '22000000-0000-0000-0000-000000000099'}, 'customer_id'),
    ({'site_id': 'missing'}, 'site_id'),
    ({'status': 'SPARE'}, 'status'),
    ({'machine_type': 'BATTERY_CABINET'}, 'machine_type'),
    ({'energy_mwh': 1}, 'energy_mwh'),
    ({'power_mw': -1}, 'power_mw'),
    ({'power_mw': '0.0000001'}, 'power_mw'),
    ({'commissioning_date': 'yesterday'}, 'commissioning_date'),
    ({'asset_id': 'existing'}, 'asset_id'),
])
def test_invalid_asset_values_produce_actionable_row_errors(database_url, change, field):
    engine, actor, imports, _ = import_context(database_url)
    rows = new_rows()
    rows[0].update(change)
    batch = imports.submit(actor, filename='invalid.xlsx', content=workbook(rows))
    imports.validate_next()
    preview = imports.get(actor, batch['id'])
    assert preview['summary']['error_rows'] >= 1
    assert any(i['row'] in (1, 2) and i['field'] == field and i['severity'] == 'ERROR' and i['message'] for i in preview['issues'])
    engine.dispose()


@pytest.mark.parametrize('variant', ['duplicate_key', 'missing_parent', 'cycle', 'earlier_child', 'duplicate_serial'])
def test_hierarchy_identity_and_serial_diagnostics(database_url, variant):
    engine, actor, imports, _ = import_context(database_url)
    rows = new_rows()
    if variant == 'duplicate_key':
        rows[1]['record_key'] = 'root'
    elif variant == 'missing_parent':
        rows[1]['parent_record_key'] = 'absent'
    elif variant == 'cycle':
        rows[0].update(location_kind='PARENT', parent_record_key='tx', customer_id=None, site_id=None)
    elif variant == 'earlier_child':
        rows[1]['effective_at'] = '2024-01-01T00:00:00Z'
    else:
        rows[1]['serial_number'] = 'TX-CURRENT-2002'
    batch = imports.submit(actor, filename='graph.xlsx', content=workbook(rows))
    imports.validate_next()
    preview = imports.get(actor, batch['id'])
    if variant == 'duplicate_serial':
        assert preview['summary']['error_count'] == 0
        assert preview['summary']['warning_rows'] == 1
        assert 'SYS-TX-002' in preview['issues'][0]['message']
    else:
        assert preview['summary']['error_count'] > 0
        assert any(i['field'] in ('record_key', 'parent_record_key', 'effective_at') for i in preview['issues'])
    engine.dispose()


@pytest.mark.parametrize('variant', ['not_zip', 'empty', 'formula', 'huge_xml', 'external_link', 'missing_header', 'too_many_rows', 'duplicate_header', 'far_column'])
def test_unsafe_or_invalid_workbook_is_retained_but_never_publishable(database_url, variant):
    engine, actor, imports, files = import_context(database_url)
    rows = new_rows()
    if variant == 'formula':
        rows[0]['serial_number'] = '=1+1'
    content = workbook(rows)
    if variant == 'not_zip':
        content = b'not an Excel workbook'
    elif variant in ('huge_xml', 'external_link'):
        stream = BytesIO(content)
        with ZipFile(stream, 'a', ZIP_DEFLATED) as archive:
            archive.writestr('xl/externalLinks/link1.xml' if variant == 'external_link' else 'huge.xml', b'x' * (11 * 1024 * 1024) if variant == 'huge_xml' else b'<x/>')
        content = stream.getvalue()
    elif variant == 'empty':
        content = workbook([])
    elif variant == 'too_many_rows':
        content = workbook(rows * 251)
    elif variant in ('missing_header', 'duplicate_header', 'far_column'):
        from openpyxl import load_workbook
        book = load_workbook(BytesIO(content))
        if variant == 'far_column':
            book['Assets']['XFD2'] = 'must not silently drop this asset field'
        else:
            book['Assets']['A1'] = 'serial_number' if variant == 'duplicate_header' else None
        stream = BytesIO()
        book.save(stream)
        content = stream.getvalue()
    batch = imports.submit(actor, filename='unsafe.xlsx', content=content)
    imports.validate_next()
    preview = imports.get(actor, batch['id'])
    assert preview['state'] == 'FAILED'
    assert preview['summary']['error_count'] > 0
    assert content in files.files.values()
    engine.dispose()


def test_upload_rejects_unauthorized_oversized_and_unavailable_storage(database_url):
    engine, actor, imports, files = import_context(database_url)
    with pytest.raises(IdentityError):
        imports.submit(replace(actor, permissions=frozenset({'assets.read', 'assets.write'})), filename='a.xlsx', content=workbook(new_rows()))
    with pytest.raises(AssetQueryError, match='import_file_size'):
        imports.submit(actor, filename='a.xlsx', content=b'x' * (2 * 1024 * 1024 + 1))
    with pytest.raises(ObjectStorageError):
        AssetImports(engine, UnavailableFileStore()).submit(actor, filename='a.xlsx', content=workbook(new_rows()))
    assert not files.files
    assert imports.validate_next() is False
    engine.dispose()
