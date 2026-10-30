"""Laboratory codec safety cases supplement the public facade acceptance tests."""
import csv
from io import BytesIO, StringIO
from zipfile import ZipFile, ZIP_DEFLATED
import pytest
from dga.laboratory.import_tables import SAMPLE_COLUMNS, TEST_COLUMNS


def csv_bytes(columns, rows, *, ending=True):
    stream=StringIO(newline='')
    writer=csv.writer(stream)
    writer.writerow(columns)
    for row in rows:
        writer.writerow([row.get(column,'') for column in columns])
    content=stream.getvalue()
    return (content if ending else content.rstrip('\r\n')).encode('utf-8')


def test_csv_pair_preserves_quoted_newline_last_record_and_complete_test_identity():
    from dga.laboratory.import_tables import read_csv_pair
    samples=csv_bytes(SAMPLE_COLUMNS,[dict(source_lab_key='LAB_A',external_sample_id='00017',
        notes='quoted, "record"\nnext line',source_note=' source evidence ')],ending=False)
    tests=csv_bytes(TEST_COLUMNS,[dict(source_lab_key='LAB_A',external_sample_id='00017',
        external_test_id='DGA-01',test_type='DGA',h2_value='0.000001',h2_qualifier='EQ')],ending=False)
    tables=read_csv_pair(samples,tests,sample_file_id='sample-file',test_file_id='test-file')
    assert [table.sheet for table in tables]==['油样信息','检测结果']
    row=tables[0].rows[0]
    assert row.values['external_sample_id']=='00017'
    assert row.values['notes']=='quoted, "record"\nnext line'
    assert row.values['source_note']==' source evidence ' and row.row_number==2
    assert tables[1].rows[0].values['external_test_id']=='DGA-01'
    assert tables[1].rows[0].values['h2_value']=='0.000001'
    assert tables[1].rows[0].cell_types['h2_value']=='TEXT'



def xlsx_bytes(samples=(),tests=()):
    from openpyxl import load_workbook
    from dga.laboratory.import_tables import build_template
    book=load_workbook(BytesIO(build_template('XLSX')['content']))
    for sheet,columns,rows in (('油样信息',SAMPLE_COLUMNS,samples),('检测结果',TEST_COLUMNS,tests)):
        for row in rows:
            book[sheet].append([row.get(column,'') for column in columns])
    stream=BytesIO()
    book.save(stream)
    return stream.getvalue()


def rewrite_zip(content, changes=None, extra=()):
    result=BytesIO()
    with ZipFile(BytesIO(content)) as source, ZipFile(result,'w',ZIP_DEFLATED) as target:
        for name in source.namelist():
            data=source.read(name)
            target.writestr(name,(changes or {}).get(name,data))
        for name,data in extra:
            target.writestr(name,data)
    return result.getvalue()


def test_xlsx_reader_preserves_numeric_lexeme_text_identity_and_boolean_type():
    from dga.laboratory.import_tables import read_xlsx_tables
    content=xlsx_bytes([dict(source_lab_key='LAB_A',external_sample_id='00017',
        sampled_at='2026-10-29',physical_received=False,source_note=' unchanged evidence ')],
        [dict(source_lab_key='LAB_A',external_sample_id='00017',external_test_id='DGA-01',
              test_type='DGA',h2_value=1.23,h2_qualifier='EQ')])
    with ZipFile(BytesIO(content)) as package:
        lexical=package.read('xl/worksheets/sheet3.xml').replace(b'<v>1.23</v>',b'<v>1.230000</v>')
    content=rewrite_zip(content,{'xl/worksheets/sheet3.xml':lexical})
    tables=read_xlsx_tables(content,file_id='workbook-file')
    assert [table.sheet for table in tables]==['油样信息','检测结果']
    assert tables[0].rows[0].values['external_sample_id']=='00017'
    assert tables[0].rows[0].values['physical_received'] is False
    assert tables[0].rows[0].cell_types['physical_received']=='BOOLEAN'
    assert tables[0].rows[0].values['source_note']==' unchanged evidence '
    assert tables[1].rows[0].values['h2_value']=='1.230000'
    assert tables[1].rows[0].cell_types['h2_value']=='NUMBER'



@pytest.mark.parametrize('content,code',[
    (b'\xff','invalid_utf8'),
    (csv_bytes(SAMPLE_COLUMNS,[]) + b'"unfinished','invalid_csv'),
    (csv_bytes(SAMPLE_COLUMNS,[]) + b'LAB_A,00"017'+b','*(len(SAMPLE_COLUMNS)-2)+b'\n','invalid_csv'),
    (csv_bytes(SAMPLE_COLUMNS,[]) + b'only,three,columns\n','schema_invalid'),
    (csv_bytes(SAMPLE_COLUMNS,[dict(source_note='bad\x00value')]),'invalid_control_character'),
    (csv_bytes(SAMPLE_COLUMNS[:-1],[]),'schema_invalid'),
    (csv_bytes(SAMPLE_COLUMNS[:-1]+('unknown',),[]),'schema_invalid'),
    (csv_bytes(SAMPLE_COLUMNS[:-1]+(SAMPLE_COLUMNS[0],),[]),'schema_invalid'),
])
def test_csv_pair_rejects_invalid_syntax_and_schema_without_partial_tables(content,code):
    from dga.laboratory.import_tables import read_csv_pair,ImportFileError
    with pytest.raises(ImportFileError,match=code):
        read_csv_pair(content,csv_bytes(TEST_COLUMNS,[]),sample_file_id='s',test_file_id='t')


def test_csv_capacity_is_complete_and_formula_text_retained_as_an_input_problem():
    from dga.laboratory.import_tables import read_csv_pair,ImportFileError
    samples=csv_bytes(SAMPLE_COLUMNS,[dict(source_lab_key='LAB_A',external_sample_id=f'{n:05}') for n in range(500)])
    tests=csv_bytes(TEST_COLUMNS,[dict(source_lab_key='LAB_A',external_sample_id='00000',
        external_test_id=f'T{n:04}',h2_value='=1+1') for n in range(2000)])
    tables=read_csv_pair(b'\xef\xbb\xbf'+samples,tests,sample_file_id='s',test_file_id='t')
    assert len(tables[0].rows)==500 and len(tables[1].rows)==2000
    assert tables[0].rows[-1].values['external_sample_id']=='00499'
    assert tables[1].rows[-1].values['external_test_id']=='T1999'
    assert tables[1].rows[-1].values['h2_value']=='=1+1'
    assert tables[1].rows[-1].problems==({'code':'formula_not_allowed','field':'h2_value','raw_value':'=1+1'},)
    with pytest.raises(ImportFileError,match='business_row_limit'):
        read_csv_pair(samples+csv_bytes(SAMPLE_COLUMNS,[dict(external_sample_id='overflow')]).split(b'\r\n',1)[1],
            tests,sample_file_id='s',test_file_id='t')
    with pytest.raises(ImportFileError,match='business_row_limit'):
        read_csv_pair(samples,tests+csv_bytes(TEST_COLUMNS,[dict(external_test_id='overflow')]).split(b'\r\n',1)[1],
            sample_file_id='s',test_file_id='t')


def test_csv_cell_limit_is_codepoints_and_does_not_truncate_or_change_global_limit():
    from dga.laboratory.import_tables import read_csv_pair,ImportFileError
    before=csv.field_size_limit()
    valid=csv_bytes(SAMPLE_COLUMNS,[dict(notes='🧪'*32767)])
    tables=read_csv_pair(valid,csv_bytes(TEST_COLUMNS,[]),sample_file_id='s',test_file_id='t')
    assert tables[0].rows[0].values['notes']=='🧪'*32767
    assert csv.field_size_limit()==before
    with pytest.raises(ImportFileError,match='cell_text_limit'):
        read_csv_pair(csv_bytes(SAMPLE_COLUMNS,[dict(notes='x'*32768)]),
            csv_bytes(TEST_COLUMNS,[]),sample_file_id='s',test_file_id='t')
    assert csv.field_size_limit()==before



@pytest.mark.parametrize('attack,code',[
    ('duplicate_member','unsafe_archive'),('path_traversal','unsafe_archive'),
    ('macro','unsafe_archive'),('external_relationship','unsafe_relationship'),
    ('formula','formula_not_allowed'),('hidden_sheet','hidden_input'),
    ('hidden_row','hidden_input'),('unknown_header','schema_invalid'),
    ('duplicate_coordinate','invalid_cell_coordinate'),('sparse_coordinate','sparse_cell_limit'),
    ('overlong_cell','cell_text_limit'),('doctype','invalid_xlsx'),
    ('wrong_workbook_root','invalid_xlsx'),('row_outside_sheetdata','invalid_xlsx'),
])
def test_xlsx_rejects_unsafe_source_before_returning_any_partial_tables(attack,code):
    from dga.laboratory.import_tables import read_xlsx_tables,ImportFileError
    content=xlsx_bytes([dict(source_lab_key='LAB_A',external_sample_id='00017',source_note='tiny')])
    with ZipFile(BytesIO(content)) as package:
        pieces={name:package.read(name) for name in package.namelist()}
    changes={}
    extra=[]
    sheet=pieces['xl/worksheets/sheet2.xml']
    if attack=='duplicate_member':extra=[('xl/workbook.xml',pieces['xl/workbook.xml'])]
    elif attack=='path_traversal':extra=[('../escape.xml',b'<x/>')]
    elif attack=='macro':extra=[('xl/vbaProject.bin',b'macro')]
    elif attack=='external_relationship':
        changes['xl/_rels/workbook.xml.rels']=pieces['xl/_rels/workbook.xml.rels'].replace(
            b'</Relationships>',b'<Relationship Id="external" Type="http://example/externalLink" Target="https://example.invalid/file" TargetMode="External"/></Relationships>')
    elif attack=='formula':changes['xl/worksheets/sheet2.xml']=sheet.replace(b'<c r="A2"',b'<c r="A2"',1).replace(b'</row>',b'<c r="S1"><f>1+1</f><v>2</v></c></row>',1)
    elif attack=='hidden_sheet':changes['xl/workbook.xml']=pieces['xl/workbook.xml'].replace(b'state="visible"',b'state="hidden"',1)
    elif attack=='hidden_row':changes['xl/worksheets/sheet2.xml']=sheet.replace(b'<row r="2"',b'<row r="2" hidden="1"',1)
    elif attack=='unknown_header':changes['xl/worksheets/sheet2.xml']=sheet.replace(b'>source_note<',b'>unknown_header<',1)
    elif attack=='duplicate_coordinate':changes['xl/worksheets/sheet2.xml']=sheet.replace(b'</row>',b'<c r="A1" t="inlineStr"><is><t>duplicate</t></is></c></row>',1)
    elif attack=='sparse_coordinate':changes['xl/worksheets/sheet2.xml']=sheet.replace(b'<row r="2"',b'<row r="1048576"',1).replace(b'r="A2"',b'r="A1048576"',1)
    elif attack=='overlong_cell':changes['xl/worksheets/sheet2.xml']=sheet.replace(b'>tiny<',b'>'+b'x'*32768+b'<',1)
    elif attack=='doctype':changes['xl/worksheets/sheet2.xml']=b'<!DOCTYPE worksheet [<!ENTITY x "value">]>'+sheet
    elif attack=='wrong_workbook_root':changes['xl/workbook.xml']=pieces['xl/workbook.xml'].replace(b'<workbook ',b'<notWorkbook ',1).replace(b'</workbook>',b'</notWorkbook>')
    elif attack=='row_outside_sheetdata':changes['xl/worksheets/sheet2.xml']=sheet.replace(b'<sheetData>',b'<notSheetData>').replace(b'</sheetData>',b'</notSheetData>')
    with pytest.raises(ImportFileError,match=code):
        read_xlsx_tables(rewrite_zip(content,changes,extra),file_id='unsafe-workbook')



def test_xlsx_shared_strings_and_excel_date_styles_preserve_type_instead_of_guessing():
    from datetime import date
    from openpyxl import load_workbook
    from dga.laboratory.import_tables import read_xlsx_tables
    content=xlsx_bytes([dict(source_lab_key='LAB_A',external_sample_id='00017')])
    book=load_workbook(BytesIO(content))
    book['油样信息']['I2']=date(2026,10,29)
    stream=BytesIO();book.save(stream);content=stream.getvalue()
    with ZipFile(BytesIO(content)) as package:
        sheet=package.read('xl/worksheets/sheet2.xml')
        relationships=package.read('xl/_rels/workbook.xml.rels')
        types=package.read('[Content_Types].xml')
    sheet=sheet.replace(b'<c r="C2" t="inlineStr"><is><t>00017</t></is></c>',b'<c r="C2" t="s"><v>0</v></c>')
    relationships=relationships.replace(b'</Relationships>',
        b'<Relationship Id="shared" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/sharedStrings" Target="/xl/sharedStrings.xml"/></Relationships>')
    types=types.replace(b'</Types>',b'<Override PartName="/xl/sharedStrings.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/></Types>')
    content=rewrite_zip(content,{'xl/worksheets/sheet2.xml':sheet,'xl/_rels/workbook.xml.rels':relationships,'[Content_Types].xml':types},
        [('xl/sharedStrings.xml',b'<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><si><t>00017</t></si></sst>')])
    tables=read_xlsx_tables(content,file_id='shared')
    row=tables[0].rows[0]
    assert row.values['external_sample_id']=='00017' and row.cell_types['external_sample_id']=='TEXT'
    assert row.cell_types['sampled_at']=='DATE_NUMBER'
    assert isinstance(row.values['sampled_at'],str) and '-' not in row.values['sampled_at']


def test_xlsx_encryption_and_corrupt_deflate_are_input_errors():
    from dga.laboratory.import_tables import read_xlsx_tables,ImportFileError
    content=xlsx_bytes()
    encrypted=bytearray(content)
    local=encrypted.index(b'PK\x03\x04');central=encrypted.index(b'PK\x01\x02')
    encrypted[local+6]|=1;encrypted[central+8]|=1
    with pytest.raises(ImportFileError,match='unsafe_archive'):
        read_xlsx_tables(bytes(encrypted),file_id='encrypted')
    with ZipFile(BytesIO(content)) as package:
        info=package.getinfo('xl/workbook.xml')
    corrupt=bytearray(content)
    at=info.header_offset
    name_len=int.from_bytes(corrupt[at+26:at+28],'little')
    extra_len=int.from_bytes(corrupt[at+28:at+30],'little')
    corrupt[at+30+name_len+extra_len]=0xff
    with pytest.raises(ImportFileError):
        read_xlsx_tables(bytes(corrupt),file_id='corrupt')



def test_xlsx_exact500_samples2000_complete_tests_and_one_sample2000_are_not_truncated():
    from dga.laboratory.import_tables import read_xlsx_tables,ImportFileError
    samples=[dict(source_lab_key='LAB_A',external_sample_id=f'{n:05}') for n in range(500)]
    tests=[]
    for n in range(2000):
        row=dict(source_lab_key='LAB_A',external_sample_id='00000',external_test_id=f'T{n:04}',test_type='DGA')
        for gas in ('h2','ch4','c2h2','c2h4','c2h6','co','co2'):
            row[gas+'_value']='0';row[gas+'_qualifier']='EQ';row[gas+'_unit']='UNKNOWN'
        tests.append(row)
    content=xlsx_bytes(samples,tests)
    tables=read_xlsx_tables(content,file_id='capacity')
    assert len(tables[0].rows)==500 and len(tables[1].rows)==2000
    assert [row.values['external_test_id'] for row in tables[1].rows]==[f'T{n:04}' for n in range(2000)]
    assert tables[0].rows[-1].values['external_sample_id']=='00499'
    assert all(row.values['external_sample_id']=='00000' for row in tables[1].rows)
    single=read_xlsx_tables(xlsx_bytes(samples[:1],tests),file_id='one-sample-2000')
    assert len(single[0].rows)==1 and len(single[1].rows)==2000
    assert single[1].rows[-1].values['external_test_id']=='T1999'
    for extra_samples,extra_tests in ((samples+[dict(external_sample_id='overflow')],tests),
                                      (samples,tests+[dict(external_test_id='overflow')])):
        with pytest.raises(ImportFileError,match='business_row_limit'):
            read_xlsx_tables(xlsx_bytes(extra_samples,extra_tests),file_id='over-capacity')


def test_xlsx_200_unique_members_is_allowed_but_201_is_rejected():
    from dga.laboratory.import_tables import read_xlsx_tables,ImportFileError
    content=xlsx_bytes()
    with ZipFile(BytesIO(content)) as package:
        remaining=200-len(package.infolist())
    extras=[(f'padding/{n}.txt',b'') for n in range(remaining)]
    assert len(read_xlsx_tables(rewrite_zip(content,extra=extras),file_id='members'))==2
    with pytest.raises(ImportFileError,match='archive_member_limit'):
        read_xlsx_tables(rewrite_zip(content,extra=extras+[('one-more.txt',b'')]),file_id='members')


def test_xlsx_expanded50MiB_boundary_is_checked_with_actual_retained_bytes():
    from dga.laboratory.import_tables import read_xlsx_tables,ImportFileError
    content=xlsx_bytes()
    with ZipFile(BytesIO(content)) as package:
        remaining=50*1024*1024-sum(info.file_size for info in package.infolist())
    assert len(read_xlsx_tables(rewrite_zip(content,extra=[('padding.txt',b'x'*remaining)]),file_id='expanded'))==2
    with pytest.raises(ImportFileError,match='archive_expanded_limit'):
        read_xlsx_tables(rewrite_zip(content,extra=[('padding.txt',b'x'*(remaining+1))]),file_id='expanded')


def instructions_xml(cells):
    return (b'<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData>'+
        b''.join(f'<row r="{n}"><c r="A{n}" t="inlineStr"><is><t>x</t></is></c></row>'.encode() for n in range(1,cells+1))+
        b'</sheetData></worksheet>')


def test_xlsx150000_actual_cells_includes_instructions_and_has_no_dimension_expansion():
    from dga.laboratory.import_tables import read_xlsx_tables,ImportFileError
    content=xlsx_bytes()
    with ZipFile(BytesIO(content)) as package:
        other=sum(package.read(name).count(b'<c ') for name in package.namelist()
            if name.startswith('xl/worksheets/') and name!='xl/worksheets/sheet4.xml')
    remaining=150000-other
    exact=rewrite_zip(content,{'xl/worksheets/sheet4.xml':instructions_xml(remaining)})
    assert len(read_xlsx_tables(exact,file_id='cells'))==2
    with pytest.raises(ImportFileError,match='cell_limit'):
        read_xlsx_tables(rewrite_zip(content,{'xl/worksheets/sheet4.xml':instructions_xml(remaining+1)}),file_id='cells')


def test_each_csv10MiB_pair20MiB_and_xlsx10MiB_accept_exact_size_but_not_more():
    from dga.laboratory.import_tables import read_csv_pair,read_xlsx_tables,ImportFileError
    from zipfile import ZIP_STORED
    limit=10*1024*1024
    samples=csv_bytes(SAMPLE_COLUMNS,[]);samples+=b'\n'*(limit-len(samples))
    tests=csv_bytes(TEST_COLUMNS,[]);tests+=b'\n'*(limit-len(tests))
    tables=read_csv_pair(samples,tests,sample_file_id='s',test_file_id='t')
    assert not tables[0].rows and not tables[1].rows
    with pytest.raises(ImportFileError,match='import_file_size'):
        read_csv_pair(samples+b'\n',tests,sample_file_id='s',test_file_id='t')
    base=xlsx_bytes()
    def stored_padding(size):
        output=BytesIO()
        with ZipFile(BytesIO(base)) as source,ZipFile(output,'w',ZIP_DEFLATED) as target:
            for name in source.namelist():target.writestr(name,source.read(name))
            target.writestr('padding.txt',b'x'*size,compress_type=ZIP_STORED)
        return output.getvalue()
    overhead=len(stored_padding(0))
    exact=stored_padding(limit-overhead)
    assert len(exact)==limit and len(read_xlsx_tables(exact,file_id='exact'))==2
    with pytest.raises(ImportFileError,match='import_file_size'):
        read_xlsx_tables(exact+b'x',file_id='over')



@pytest.mark.parametrize('suffix',['XML','dat'])
def test_referenced_xml_parts_cannot_bypass_formula_checks_by_filename_extension(suffix):
    from dga.laboratory.import_tables import read_xlsx_tables,ImportFileError
    original=xlsx_bytes()
    result=BytesIO()
    old='xl/worksheets/sheet4.xml'
    new='xl/worksheets/sheet4.'+suffix
    formula=b'<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData><row r="1"><c r="A1"><f>1+1</f><v>2</v></c></row></sheetData></worksheet>'
    with ZipFile(BytesIO(original)) as source,ZipFile(result,'w',ZIP_DEFLATED) as target:
        for name in source.namelist():
            data=source.read(name)
            if name==old:
                target.writestr(new,formula)
            else:
                if name in ('xl/_rels/workbook.xml.rels','[Content_Types].xml'):
                    data=data.replace(old.encode(),new.encode())
                target.writestr(name,data)
    with pytest.raises(ImportFileError,match='formula_not_allowed'):
        read_xlsx_tables(result.getvalue(),file_id='extension-bypass')
