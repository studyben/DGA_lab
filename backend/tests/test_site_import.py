"""Customer/site-only import behavior at the approved assets public seams."""
from datetime import datetime, timezone
from io import BytesIO
from uuid import uuid4
from decimal import Decimal

from openpyxl import Workbook, load_workbook
from dga.assets.public import AssetImports, AssetDirectory
from tests.test_sample_reception import reception_context, CUSTOMER_ID, SITE_ID


class Files:
    def __init__(self): self.files = {}
    def put(self, *, object_key, content, content_type): self.files[object_key] = content
    def get(self, *, object_key): return self.files[object_key]
    def delete(self, *, object_key): del self.files[object_key]


def document(sites=None, customers=None):
    book = Workbook()
    book.active.title = 'Format'
    book.active.append(['version', 'mode'])
    book.active.append([2, 'SITES'])
    defaults = {
        'Customers': ['record_key','decision','customer_id','customer_name','source_note','reason'],
        'Sites': ['record_key','decision','site_id','site_name','customer_ref','location_text','product_lines','pv_power_mw','ess_power_mw','ess_energy_mwh','commissioning_date','grid_status','operation_status','source_note','reason'],
    }
    for name, rows in [('Customers', customers or []), ('Sites', sites if sites is not None else [new_site()])]:
        sheet = book.create_sheet(name)
        headers = defaults[name] + sorted({key for r in rows for key in r} - set(defaults[name]))
        sheet.append(headers)
        for row in rows:
            sheet.append([row.get(key) for key in headers])
            for cell in sheet[sheet.max_row]:
                if isinstance(cell.value, str): cell.data_type = 's'
    output = BytesIO(); book.save(output)
    return output.getvalue()


def new_site(**changes):
    return dict(record_key='site-01', decision='NEW', site_name='New dual line site',
        customer_ref=str(CUSTOMER_ID), product_lines='PV;ESS', pv_power_mw='12.000001',
        ess_power_mw='20', ess_energy_mwh='50', source_note='Serial source: service report; capacity: shipping records',
        reason='Confirmed a separate project') | changes


def context(database_url):
    engine, identity, actor = reception_context(database_url)
    files = Files()
    return engine, identity, actor, files, AssetImports(engine, files, clock=lambda: datetime(2026,10,29,tzinfo=timezone.utc))


def upload(imports, actor, content=None, **changes):
    return imports.submit(actor, filename='site-import.xlsx', content=content or document(),
        format_version=2, mode='SITES', idempotency_key=str(uuid4()), **changes)


def test_only_site_batch_creates_customer_and_zero_device_site_with_retained_provenance(database_url):
    engine, _, actor, files, imports = context(database_url)
    customers = [dict(record_key='c-01',decision='NEW',customer_name='New customer',source_note='Customer contract',reason='Confirmed new customer')]
    batch = upload(imports, actor, document([new_site(customer_ref='NEW:c-01')], customers))
    assert batch['state'] == 'STAGED'
    assert AssetDirectory(engine).dashboard(actor, product_line='ESS')['site_count'] == 0
    assert imports.validate_next()
    preview = imports.get(actor,batch['id'])
    assert preview['summary']['error_count'] == 0, preview['issues']
    operation = str(uuid4())
    result = imports.publish(actor,batch['id'],validation_revision=preview['validation_revision'],
        idempotency_key=operation,acknowledged_warning_ids=[])
    assert result['state'] == 'PUBLISHED'
    site_id = next(r['target_id'] for r in result['rows'] if r['sheet']=='Sites')
    directory = AssetDirectory(engine)
    pv = directory.site_detail(actor,site_id,product_line='PV')
    ess = directory.site_detail(actor,site_id,product_line='ESS')
    assert pv['equipment_count'] == ess['equipment_count'] == 0
    assert pv['site']['customer_name'] == 'New customer'
    assert pv['site']['power_mw'] == Decimal('12.000001')
    assert ess['site']['power_mw'] == Decimal('20')
    assert ess['site']['energy_mwh'] == Decimal('50')
    source = directory.import_sources(actor,'SITE',site_id)['sources'][0]
    assert source['source_note'] == new_site()['source_note']
    assert source['original_values']['customer_ref'] == 'NEW:c-01'
    assert imports.source_file(actor,batch['id'])['content'] in files.files.values()
    again = imports.publish(actor,batch['id'],validation_revision=preview['validation_revision'],
        idempotency_key=operation,acknowledged_warning_ids=[])
    assert [r['target_id'] for r in again['rows']] == [r['target_id'] for r in result['rows']]
    assert imports.operation(actor,operation)['outcome'] == 'PUBLISHED'
    assert site_id in imports.operation(actor,operation)['target_ids']
    exported = directory.export_sites(actor)
    try:
        book = load_workbook(exported.file,read_only=True)
        rows = list(book['Sites'].values)
        imported = next(dict(zip(rows[0],r)) for r in rows[1:] if r[0] == site_id)
        assert imported['source_note'] == new_site()['source_note']
    finally: exported.file.close()
    engine.dispose()


def test_same_name_reference_requires_review_and_keeps_existing_customer_and_capacity(database_url):
    import pytest
    from sqlalchemy import text
    from dga.assets.public import AssetQueryError
    engine, _, actor, _, imports = context(database_url)
    with engine.begin() as c:
        c.execute(text("INSERT INTO site_product_lines(site_id,product_line,power_mw) VALUES (:id,'PV',7)"),{'id':SITE_ID})
    batch=upload(imports,actor,document([new_site(site_name='Prairie Sun',product_lines='PV',pv_power_mw='99',ess_power_mw=None,ess_energy_mwh=None)]))
    imports.validate_next(); preview=imports.get(actor,batch['id'])
    assert preview['summary']['warning_count']>0
    revised=imports.revise(actor,batch['id'],expected_revision=preview['validation_revision'],
        changes=[dict(sheet='Sites',row=2,values=dict(decision='REFERENCE',site_id=str(SITE_ID),reason='Confirmed existing site'))],
        reason='Use official site identity',idempotency_key=str(uuid4()))
    assert revised['validation_revision']>preview['validation_revision']
    history=imports.revision_history(actor,batch['id'])
    assert history['changes'][0]['before']['decision']=='NEW'
    assert history['changes'][0]['after']['decision']=='REFERENCE'
    assert any(i['code']=='reference_difference' for i in revised['issues'])
    with pytest.raises(AssetQueryError) as error:
        imports.publish(actor,batch['id'],validation_revision=revised['validation_revision'],idempotency_key=str(uuid4()),acknowledged_warning_ids=[])
    assert error.value.code=='import_warnings_unacknowledged'
    result=imports.publish(actor,batch['id'],validation_revision=revised['validation_revision'],idempotency_key=str(uuid4()),acknowledged_warning_ids=[i['id'] for i in revised['issues'] if i['severity']=='WARNING'])
    assert result['state']=='PUBLISHED'
    detail=AssetDirectory(engine).site_detail(actor,SITE_ID,product_line='PV')
    assert detail['site']['power_mw']==7
    assert detail['site']['customer_id']==CUSTOMER_ID
    source=AssetDirectory(engine).import_sources(actor,'SITE',SITE_ID)['sources'][0]
    assert source['original_values']['decision']=='NEW'
    assert source['reviewed_values']['decision']=='REFERENCE'
    assert source['reviewed_values']['pv_power_mw']=='99'
    assert detail['site']['source_note'] is None
    engine.dispose()


def test_preview_catalog_change_requires_new_revision_even_with_same_publish_intent(database_url):
    from sqlalchemy import text
    engine, _, actor, _, imports=context(database_url)
    batch=upload(imports,actor); imports.validate_next(); preview=imports.get(actor,batch['id'])
    with engine.begin() as c:
        c.execute(text("INSERT INTO sites(id,customer_id,site_name) VALUES (:id,:c,'New dual line site')"),dict(id=uuid4(),c=CUSTOMER_ID))
    key=str(uuid4())
    refreshed=imports.publish(actor,batch['id'],validation_revision=preview['validation_revision'],idempotency_key=key,acknowledged_warning_ids=[])
    assert refreshed['state']=='VALIDATED'
    assert refreshed['validation_revision']>preview['validation_revision']
    assert imports.operation(actor,key)['outcome']=='REFRESHED'
    replay=imports.publish(actor,batch['id'],validation_revision=preview['validation_revision'],idempotency_key=key,acknowledged_warning_ids=[])
    assert replay['state']=='VALIDATED'
    assert not any(r.get('target_id') for r in replay['rows'])
    engine.dispose()


def test_publish_rejects_missing_or_corrupted_source_without_creating_site(database_url):
    import pytest
    from dga.shared.files import ObjectStorageError
    engine, _, actor, files, imports=context(database_url)
    batch=upload(imports,actor); imports.validate_next(); preview=imports.get(actor,batch['id'])
    original=dict(files.files)
    for replacement in ({},{next(iter(original)):b'corrupted'}):
        files.files=replacement
        with pytest.raises(ObjectStorageError):
            imports.publish(actor,batch['id'],validation_revision=preview['validation_revision'],idempotency_key=str(uuid4()),acknowledged_warning_ids=[])
        assert AssetDirectory(engine).dashboard(actor,product_line='ESS')['site_count']==0
        assert imports.get(actor,batch['id'])['state']=='VALIDATED'
    files.files=original
    engine.dispose()


def test_concurrent_same_intent_and_lost_response_can_resolve_persistent_operation(database_url):
    import pytest
    from concurrent.futures import ThreadPoolExecutor
    from dga.assets.public import AssetQueryError
    engine, _, actor, _, imports=context(database_url)
    content=document(); key=str(uuid4())
    def submit(): return imports.submit(actor,filename='site-import.xlsx',content=content,format_version=2,mode='SITES',idempotency_key=key)
    with ThreadPoolExecutor(max_workers=2) as pool: batches=list(pool.map(lambda _:submit(),range(2)))
    assert batches[0]['id']==batches[1]['id']
    with pytest.raises(AssetQueryError) as error:
        imports.submit(actor,filename='changed.xlsx',content=content,format_version=2,mode='SITES',idempotency_key=key)
    assert error.value.code=='import_operation_conflict'
    imports.validate_next(); preview=imports.get(actor,batches[0]['id']); publish_key=str(uuid4())
    def publish(): return imports.publish(actor,batches[0]['id'],validation_revision=preview['validation_revision'],idempotency_key=publish_key,acknowledged_warning_ids=[])
    with ThreadPoolExecutor(max_workers=2) as pool: results=list(pool.map(lambda _:publish(),range(2)))
    assert results[0]['rows']==results[1]['rows']
    assert imports.operation(actor,publish_key)['outcome']=='PUBLISHED'
    assert AssetDirectory(engine).dashboard(actor,product_line='ESS')['site_count']==1
    engine.dispose()


def test_site_publication_rolls_back_late_failure_with_source_and_retry_intent_retained(database_url):
    import pytest
    from sqlalchemy import event
    engine, _, actor, files, imports=context(database_url)
    batch=upload(imports,actor); imports.validate_next(); preview=imports.get(actor,batch['id']); key=str(uuid4())
    def fail_database_audit(connection,cursor,statement,parameters,context,many):
        if 'INSERT INTO audit_logs' in statement and 'SITE_IMPORT_PUBLISH' in str(parameters):
            raise RuntimeError('injected database transport failure at final audit')
    event.listen(engine,'before_cursor_execute',fail_database_audit)
    try:
        with pytest.raises(RuntimeError): imports.publish(actor,batch['id'],validation_revision=preview['validation_revision'],idempotency_key=key,acknowledged_warning_ids=[])
    finally: event.remove(engine,'before_cursor_execute',fail_database_audit)
    assert AssetDirectory(engine).dashboard(actor,product_line='ESS')['site_count']==0
    assert imports.get(actor,batch['id'])['state']=='VALIDATED'
    assert imports.source_file(actor,batch['id'])['content'] in files.files.values()
    assert imports.publish(actor,batch['id'],validation_revision=preview['validation_revision'],idempotency_key=key,acknowledged_warning_ids=[])['state']=='PUBLISHED'
    engine.dispose()


def test_reuploaded_file_links_previous_batch_and_requires_explicit_duplicate_review(database_url):
    engine, _, actor, _, imports=context(database_url)
    content=document(); first=upload(imports,actor,content)
    second=upload(imports,actor,content); imports.validate_next(); imports.validate_next()
    preview=imports.get(actor,second['id'])
    assert str(first['id']) in [b['id'] for b in preview['duplicate_batches']]
    assert any(i['code']=='duplicate_file' for i in preview['issues'])
    engine.dispose()


import pytest

@pytest.mark.parametrize('change,field',[
    ({'pv_power_mw':'0.0000001'},'pv_power_mw'),({'pv_power_mw':'NaN'},'pv_power_mw'),
    ({'ess_energy_mwh':'-1'},'ess_energy_mwh'),({'source_note':'x'*2001},'source_note'),
    ({'site_name':123},'site_name'),({'record_key':123},'record_key'),
    ({'decision':'UPDATE'},'decision'),({'customer_ref':'missing'},'customer_ref'),
    ({'product_lines':'PV;BAD'},'product_lines'),({'operation_status':'MAINTENANCE'},'operation_status'),
    ({'commissioning_date':'2026-02-30'},'commissioning_date'),({'grid_year':2020},'headers'),
])
def test_invalid_customer_site_values_are_actionable_and_never_publishable(database_url,change,field):
    engine, _, actor, _, imports=context(database_url)
    batch=upload(imports,actor,document([new_site(**change)]))
    if batch['state']=='STAGED': imports.validate_next()
    preview=imports.get(actor,batch['id'])
    assert preview['summary']['error_count']>0
    assert any(i['field']==field and i['message'] for i in preview['issues'])
    assert AssetDirectory(engine).dashboard(actor,product_line='ESS')['site_count']==0
    engine.dispose()


def test_current_grants_are_rechecked_and_lab_manager_import_is_additive(database_url):
    from tests.identity_seed import seed_legacy_user
    from tests.test_sample_reception import INITIAL_PASSWORD,CHANGED_PASSWORD
    from dga.shared.auth.public import IdentityError
    engine, identity, admin, _, imports=context(database_url)
    manager_id=seed_legacy_user(database_url,'site-manager','Site manager',INITIAL_PASSWORD,['lab_admin'])
    first=identity.login('site-manager',INITIAL_PASSWORD)
    manager=identity.change_password(first.token,INITIAL_PASSWORD,CHANGED_PASSWORD).actor
    assert {'assets.import','assets.write','assets.site.edit'}<=manager.permissions
    batch=upload(imports,manager); imports.validate_next(); preview=imports.get(manager,batch['id'])
    identity.change_roles(admin,manager_id,add=['field_engineer'],remove=['lab_admin'],expected_revision=0)
    with pytest.raises(IdentityError):
        imports.publish(manager,batch['id'],validation_revision=preview['validation_revision'],idempotency_key=str(uuid4()),acknowledged_warning_ids=[])
    assert AssetDirectory(engine).dashboard(admin,product_line='ESS')['site_count']==0
    engine.dispose()


def test_lost_upload_commit_response_is_resolved_without_deleting_source(database_url):
    from contextlib import contextmanager
    engine, _, actor, files, imports=context(database_url)
    class LostResponse:
        def __init__(self): self.armed=True
        def connect(self): return engine.connect()
        @contextmanager
        def begin(self):
            with engine.begin() as c: yield c
            if self.armed:
                self.armed=False
                raise ConnectionError('COMMIT response lost after durable commit')
    failed=AssetImports(LostResponse(),files)
    key=str(uuid4()); content=document()
    with pytest.raises(ConnectionError):
        failed.submit(actor,filename='site.xlsx',content=content,format_version=2,mode='SITES',idempotency_key=key)
    outcome=imports.operation(actor,key)
    assert outcome['outcome']=='STAGED'
    assert imports.source_file(actor,outcome['batch_id'])['content']==content
    again=imports.submit(actor,filename='site.xlsx',content=content,format_version=2,mode='SITES',idempotency_key=key)
    assert str(again['id'])==outcome['batch_id']
    engine.dispose()


def test_500_site_rows_are_processed_without_paging_truncation_and_501_is_rejected(database_url):
    engine, _, actor, _, imports=context(database_url)
    sites=[new_site(record_key=f's-{n}',site_name=f'Capacity boundary {n}') for n in range(500)]
    batch=upload(imports,actor,document(sites)); imports.validate_next(); preview=imports.get(actor,batch['id'])
    assert preview['total_rows']==500
    assert len(preview['rows'])==20
    assert preview['summary']['error_count']==0
    result=imports.publish(actor,batch['id'],validation_revision=preview['validation_revision'],idempotency_key=str(uuid4()),acknowledged_warning_ids=[])
    assert result['state']=='PUBLISHED'
    assert AssetDirectory(engine).dashboard(actor,product_line='ESS')['site_count']==500
    assert len(imports.get(actor,batch['id'],page=25)['rows'])==20
    too_many=upload(imports,actor,document(sites+[new_site(record_key='overflow')]))
    assert too_many['state']=='FAILED'
    assert '500' in too_many['issues'][0]['message']
    engine.dispose()


@pytest.mark.parametrize('variant',['formula','external','duplicate_zip','far_row','empty_sites','equipment_rows'])
def test_unsafe_v2_workbooks_are_retained_with_visible_errors(database_url,variant):
    from zipfile import ZipFile,ZIP_DEFLATED
    engine, _, actor, _, imports=context(database_url)
    content=document()
    if variant in ('formula','far_row','empty_sites','equipment_rows'):
        book=load_workbook(BytesIO(content))
        if variant=='formula': book['Sites']['D2']='=1+1'
        elif variant=='far_row': book['Sites']['A200000']='far-away-content'
        elif variant=='empty_sites': book['Sites'].delete_rows(2)
        else:
            tab=book.create_sheet('Assets');tab.append(['serial_number']);tab.append(['unintended-write'])
        output=BytesIO();book.save(output);content=output.getvalue()
    else:
        output=BytesIO(content)
        with ZipFile(output,'a',ZIP_DEFLATED) as archive:
            archive.writestr('xl/externalLinks/extra.xml' if variant=='external' else 'xl/workbook.xml',b'<x/>')
        content=output.getvalue()
    batch=upload(imports,actor,content)
    assert batch['state']=='FAILED'
    assert batch['issues'] and batch['issues'][0]['message']
    assert imports.source_file(actor,batch['id'])['content']==content
    engine.dispose()


@pytest.mark.parametrize('database_url',['0030_merge_export_identity'],indirect=True)
def test_real_0030_upgrade_retains_v1_source_and_reference_sheet_semantics(database_url,monkeypatch):
    from sqlalchemy import text
    from alembic import command
    from alembic.config import Config
    from tests.test_asset_import import workbook,new_rows
    engine, _, actor = reception_context(database_url)
    content=workbook(new_rows()); batch_id=uuid4(); files=Files()
    files.put(object_key='legacy/source.xlsx',content=content,content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    import hashlib
    with engine.begin() as c:
        c.execute(text("""INSERT INTO asset_import_batches(id,filename,sha256,object_key,byte_size,state,source_rows,created_by,created_at)
            VALUES (:id,'v1.xlsx',:sha,'legacy/source.xlsx',:size,'STAGED','[]',:actor,now())"""),dict(id=batch_id,sha=hashlib.sha256(content).hexdigest(),size=len(content),actor=actor.user_id))
    monkeypatch.setenv('DATABASE_URL',database_url)
    command.upgrade(Config('alembic.ini'),'head')
    imports=AssetImports(engine,files)
    result=imports.get(actor,batch_id)
    assert result['format_version']==1 and result['mode']=='ASSETS'
    assert result['state']=='STAGED'
    assert imports.source_file(actor,batch_id)['content']==content
    assert AssetDirectory(engine).dashboard(actor,product_line='ESS')['site_count']==0
    v1=load_workbook(BytesIO(imports.template(actor)),read_only=True)
    assert 'Assets' in v1.sheetnames and 'Format' not in v1.sheetnames
    engine.dispose()


def test_500_customers_and_500_sites_remain_one_atomic_batch(database_url):
    engine, _, actor, _, imports=context(database_url)
    customers=[dict(record_key=f'c-{n}',decision='NEW',customer_name=f'Bulk customer {n}',reason='Verified new customer',source_note='x'*2000) for n in range(500)]
    sites=[new_site(record_key=f's-{n}',site_name=f'Bulk site {n}',customer_ref=f'NEW:c-{n}') for n in range(500)]
    batch=upload(imports,actor,document(sites,customers));imports.validate_next();preview=imports.get(actor,batch['id'])
    assert preview['total_rows']==1000 and preview['summary']['error_count']==0
    result=imports.publish(actor,batch['id'],validation_revision=preview['validation_revision'],idempotency_key=str(uuid4()),acknowledged_warning_ids=[])
    assert result['state']=='PUBLISHED'
    totals=AssetDirectory(engine).dashboard(actor,product_line='ESS')
    assert totals['site_count']==500 and totals['customer_count']==500
    assert imports.get(actor,batch['id'],page=50)['rows'][-1]['target_id']
    engine.dispose()
