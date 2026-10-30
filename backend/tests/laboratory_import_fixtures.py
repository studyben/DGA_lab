"""Synthetic complete-row inputs; never a production or accepted scientific catalog."""
import csv
from io import StringIO,BytesIO
from uuid import uuid4
from openpyxl import load_workbook
from dga.laboratory.public import LaboratoryImports,LaboratoryConfiguration,MethodVersionInput
from tests.test_sample_reception import reception_context,TRANSFORMER_ID
from tests.test_laboratory_imports import Files

SAMPLES='source_lab_key,source_lab_name,external_sample_id,source_site_name,source_whole_unit_serial,source_transformer_serial,formal_asset_id,system_asset_number,sampled_at,sampled_timezone,sampling_context,physical_received,received_at,received_timezone,container_count,report_filenames,notes,source_note'.split(',')
TESTS='source_lab_key,external_sample_id,external_test_id,test_type,measured_at,measured_timezone,source_method,source_method_version,method_version_id,report_selected,source_instrument,source_analyst,source_qa,h2_value,h2_qualifier,h2_unit,ch4_value,ch4_qualifier,ch4_unit,c2h2_value,c2h2_qualifier,c2h2_unit,c2h4_value,c2h4_qualifier,c2h4_unit,c2h6_value,c2h6_qualifier,c2h6_unit,co_value,co_qualifier,co_unit,co2_value,co2_qualifier,co2_unit,result_value,result_qualifier,result_unit,notes,source_note'.split(',')


def csv_bytes(columns,rows):
    out=StringIO(newline='')
    writer=csv.writer(out)
    writer.writerow(columns)
    writer.writerows([[row.get(c,'') for c in columns] for row in rows])
    return out.getvalue().encode('utf-8')


def setup(database_url,*,format='CSV_PAIR',files=None,clock=None):
    engine,identity,actor=reception_context(database_url)
    files=files or Files()
    imports=LaboratoryImports(engine,files,clock=clock)
    config=LaboratoryConfiguration(engine)
    method=config.create_method(actor,MethodVersionInput(test_type='MOISTURE',display_name='Synthetic supplied method',
        version_label='S2-'+uuid4().hex[:12],fields=[dict(code='MOISTURE',display_name='Source value',
            unit_code='TEST-UNIT',display_decimal_places=6,minimum='0',maximum='100')]))
    name='Explicit source '+uuid4().hex
    choice=dict(key='LAB_A',decision='NEW',proposed_name=name,source_kind='THIRD_PARTY',
        reason='Synthetic source explicitly confirmed',confirm_distinct=True,
        candidate_digest=imports.source_labs(actor,exact_name=name)['candidate_digest'])
    batch=imports.create(actor,idempotency_key=str(uuid4()),input_format=format,purpose='HISTORICAL',
        source_labs=[choice])['batch']
    sample=dict(source_lab_key='LAB_A',source_lab_name=name,external_sample_id='00017',
        formal_asset_id=str(TRANSFORMER_ID),sampled_at='2026-01-15',sampling_context='OTHER',
        physical_received='FALSE',container_count='0',notes='ordinary notes',
        source_note=' source evidence,\n原始来源 ')
    test=dict(source_lab_key='LAB_A',external_sample_id='00017',external_test_id='T-0001',
        test_type='MOISTURE',measured_at='2026-01-16',source_method='Supplied source method',
        method_version_id=str(method['id']),source_instrument='Source instrument',
        source_analyst='Original source scientist',source_qa='Source evidence only',
        result_value='1.230000',result_qualifier='EQ',result_unit='TEST-UNIT',
        source_note='test evidence',notes='ordinary test notes')
    return engine,identity,actor,imports,files,batch,sample,test,method


def receive(imports,actor,batch,samples,tests):
    if batch['input_format']=='CSV_PAIR':
        data=[('SAMPLES','samples.csv','text/csv',csv_bytes(SAMPLES,samples)),
              ('TESTS','tests.csv','text/csv',csv_bytes(TESTS,tests))]
    else:
        book=load_workbook(BytesIO(imports.template(actor)['content']))
        for name,columns,rows in [('油样信息',SAMPLES,samples),('检测结果',TESTS,tests)]:
            for row in rows:
                book[name].append([row.get(c,'') for c in columns])
                for cell in book[name][book[name].max_row]:
                    if isinstance(cell.value,str): cell.data_type='s'
        out=BytesIO();book.save(out)
        data=[('WORKBOOK','source.xlsx','application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',out.getvalue())]
    result=[]
    for role,name,mime,content in data:
        result.append(imports.receive_file(actor,batch['id'],idempotency_key=str(uuid4()),role=role,
            filename=name,content_type=mime,content=content))
    return result


def queued(imports,actor,batch):
    current=imports.get(actor,batch['id'])
    return imports.validate(actor,batch['id'],idempotency_key=str(uuid4()),expected_revision=current['revision'])



def dga_method(engine,actor):
    return LaboratoryConfiguration(engine).create_method(actor,MethodVersionInput(test_type='DGA',
        display_name='Synthetic complete DGA',version_label='FULL-'+uuid4().hex[:12],
        fields=[dict(code=c,display_name=c,unit_code='TEST-UNIT',display_decimal_places=6)
                for c in ('H2','CH4','C2H2','C2H4','C2H6','CO','CO2')]))


def dga_row(test,method):
    row={k:v for k,v in test.items() if not k.startswith('result_')}
    row.update(test_type='DGA',method_version_id=str(method['id']))
    for code in ('h2','ch4','c2h2','c2h4','c2h6','co','co2'):
        row.update({code+'_value':'0.000001',code+'_qualifier':'EQ',code+'_unit':'TEST-UNIT'})
    return row

