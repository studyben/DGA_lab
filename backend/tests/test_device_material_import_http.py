"""Transport/security assertions; domain assertions live at the asset public facade."""
from io import BytesIO
from uuid import uuid4
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from dga.main import create_app
from dga.shared.config import Settings
from tests.test_device_material_import import context, document
from tests.test_sample_reception import CHANGED_PASSWORD


def test_default_v2_template_and_explicit_v1_upload_remain_distinct(database_url):
    engine, actor, imports, files = context(database_url)
    app = create_app(Settings(database_url=database_url, cookie_secure=False, auth_allowed_origins='http://testserver'), file_store=files)
    with TestClient(app) as client:
        assert client.get('/api/assets/imports/template').status_code == 401
        client.post('/api/auth/login', headers={'Origin': 'http://testserver'}, json=dict(username='reception-admin', password=CHANGED_PASSWORD))
        current = client.get('/api/assets/imports/template')
        assert current.status_code == 200
        book = load_workbook(BytesIO(current.content))
        assert list(book['Format'].values) == [('version', 'mode'), (2, 'ASSETS')]
        assert 'customer_id' not in next(book['Assets'].values)
        legacy = load_workbook(BytesIO(client.get('/api/assets/imports/template?format_version=1&mode=ASSETS').content))
        assert 'Format' not in legacy.sheetnames
        assert 'customer_id' in next(legacy['Assets'].values)
    engine.dispose()


def test_v2_assets_http_upload_revision_and_recovery_require_current_permission_and_csrf(database_url):
    engine, actor, imports, files = context(database_url)
    app = create_app(Settings(database_url=database_url, cookie_secure=False, auth_allowed_origins='http://testserver'), file_store=files)
    with TestClient(app) as client:
        route = '/api/assets/imports'
        login = client.post('/api/auth/login', headers={'Origin': 'http://testserver'}, json=dict(username='reception-admin', password=CHANGED_PASSWORD)).json()
        headers = {'Origin': 'http://testserver', 'X-CSRF-Token': login['csrf_token']}
        target = route + '?format_version=2&mode=ASSETS&filename=devices.xlsx'
        mime = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        content = document()
        assert client.post(target, content=content, headers={'Content-Type': mime}).status_code == 403
        assert client.post(target, content=content, headers=headers | {'Content-Type': 'text/plain'}).status_code == 415
        key = str(uuid4())
        uploaded = client.post(target, content=content, headers=headers | {'Content-Type': mime, 'Idempotency-Key': key})
        assert uploaded.status_code == 201, uploaded.text
        batch = uploaded.json()
        assert batch['state'] == 'STAGED' and batch['operation_result']['outcome'] == 'STAGED'
        assert client.get(route + '/operations/' + key).json() == batch['operation_result']
        assert imports.validate_next()
        preview = client.get(route + '/' + batch['id'] + '?duplicate_page=1').json()
        revised = client.post(route + '/' + batch['id'] + '/revisions', headers=headers, json=dict(
            expected_revision=preview['validation_revision'], changes=[dict(sheet='Assets', row=2, values=dict(source_note='HTTP reviewed evidence'))], reason='Checked source', idempotency_key=str(uuid4())))
        assert revised.status_code == 200, revised.text
        assert revised.json()['operation_result']['outcome'] == 'REVISED'
        assert client.get(route + '/' + batch['id'] + '/revisions').json()['changes'][0]['after']['source_note'] == 'HTTP reviewed evidence'
        bad = client.post(route + '/' + batch['id'] + '/revisions', headers=headers, json=dict(
            expected_revision=revised.json()['validation_revision'], changes=[dict(sheet='Assets', row=2, values=dict(customer_id=str(uuid4())))], reason='Forbidden', idempotency_key=str(uuid4())))
        assert bad.status_code == 422
        published = client.post(route + '/' + batch['id'] + '/publish', headers=headers, json=dict(validation_revision=revised.json()['validation_revision'], idempotency_key=str(uuid4()), acknowledged_warning_ids=[]))
        assert published.status_code == 409 and published.json() == {'code': 'import_mode_not_publishable'}
        original = client.get(route + '/' + batch['id'] + '/source')
        assert original.content == content and original.headers['cache-control'] == 'no-store'
        assert client.get(route + '/' + batch['id'] + '?duplicate_page=0').status_code == 422
        assert 'object_key' not in uploaded.text and 'catalog_digest' not in uploaded.text
    engine.dispose()
