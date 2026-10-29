"""Only transport/security assertions; business tests use public application seams."""
from io import BytesIO
from uuid import uuid4
from openpyxl import load_workbook
from fastapi.testclient import TestClient
from dga.main import create_app
from dga.shared.config import Settings
from tests.test_site_import import context,document,new_site
from tests.test_sample_reception import CHANGED_PASSWORD


def test_v2_http_template_revision_source_and_publication_keep_csrf_and_source_permissions(database_url):
    engine, _, actor, files, imports=context(database_url)
    app=create_app(Settings(database_url=database_url,cookie_secure=False,auth_allowed_origins='http://testserver'),file_store=files)
    with TestClient(app) as client:
        route='/api/assets/imports'
        assert client.get(route+'/template?format_version=2&mode=SITES').status_code==401
        login=client.post('/api/auth/login',headers={'Origin':'http://testserver'},json=dict(username='reception-admin',password=CHANGED_PASSWORD)).json()
        headers={'Origin':'http://testserver','X-CSRF-Token':login['csrf_token']}
        template=client.get(route+'/template?format_version=2&mode=SITES')
        assert template.status_code==200
        assert 'Format' in load_workbook(BytesIO(template.content)).sheetnames
        target=route+'?format_version=2&mode=SITES&filename=site.xlsx'
        assert client.post(target,content=document(),headers={'Content-Type':template.headers['content-type']}).status_code==403
        key=str(uuid4())
        uploaded=client.post(target,content=document(),headers=headers|{'Content-Type':template.headers['content-type'],'Idempotency-Key':key})
        assert uploaded.status_code==201,uploaded.text
        batch_id=uploaded.json()['id'];imports.validate_next()
        preview=client.get(route+'/'+batch_id).json()
        bad=client.post(route+'/'+batch_id+'/revisions',headers=headers,json=dict(expected_revision=preview['validation_revision'],changes=[],reason='bad',idempotency_key=str(uuid4()),unexpected=True))
        assert bad.status_code==422
        revised=client.post(route+'/'+batch_id+'/revisions',headers=headers,json=dict(expected_revision=preview['validation_revision'],changes=[dict(sheet='Sites',row=2,values=dict(source_note='Rechecked shipping records'))],reason='Evidence checked',idempotency_key=str(uuid4())))
        assert revised.status_code==200,revised.text
        assert client.get(route+'/'+batch_id+'/revisions').json()['changes'][0]['after']['source_note']=='Rechecked shipping records'
        source=client.get(route+'/'+batch_id+'/source')
        assert source.status_code==200 and source.headers['cache-control']=='no-store'
        assert client.get(route+'/operations/'+key).json()['batch_id']==batch_id
        published=client.post(route+'/'+batch_id+'/publish',headers=headers,json=dict(validation_revision=revised.json()['validation_revision'],idempotency_key=str(uuid4()),acknowledged_warning_ids=[]))
        assert published.status_code==200,published.text
        site_id=published.json()['rows'][0]['target_id']
        assert client.get('/api/assets/sites/'+site_id+'/sources').json()['sources'][0]['source_note']=='Rechecked shipping records'
    engine.dispose()


def test_asset_reader_sees_provenance_but_cannot_download_private_import_source(database_url):
    from tests.identity_seed import seed_legacy_user
    from tests.test_sample_reception import INITIAL_PASSWORD
    engine,identity,admin,files,imports=context(database_url)
    batch=imports.submit(admin,filename='permissions.xlsx',content=document(),format_version=2,mode='SITES',idempotency_key=str(uuid4()))
    imports.validate_next();preview=imports.get(admin,batch['id'])
    result=imports.publish(admin,batch['id'],validation_revision=preview['validation_revision'],idempotency_key=str(uuid4()),acknowledged_warning_ids=[])
    site_id=result['rows'][0]['target_id']
    seed_legacy_user(database_url,'source-reader','Reader',INITIAL_PASSWORD,['field_engineer'])
    login=identity.login('source-reader',INITIAL_PASSWORD);identity.change_password(login.token,INITIAL_PASSWORD,CHANGED_PASSWORD)
    with TestClient(create_app(Settings(database_url=database_url,cookie_secure=False,auth_allowed_origins='http://testserver'),file_store=files)) as client:
        assert client.get('/api/assets/sites/'+site_id+'/sources').status_code==401
        logged=client.post('/api/auth/login',headers={'Origin':'http://testserver'},json=dict(username='source-reader',password=CHANGED_PASSWORD))
        assert logged.status_code==200
        source=client.get('/api/assets/sites/'+site_id+'/sources')
        assert source.status_code==200
        assert source.json()['sources'][0]['source_note']==new_site()['source_note']
        assert 'object_key' not in source.text
        assert client.get('/api/assets/imports/'+str(batch['id'])+'/source').status_code==403
        assert client.get('/api/assets/imports/'+str(batch['id'])).status_code==403
        assert client.get('/api/assets/imports/template?format_version=2&mode=SITES').status_code==403
    engine.dispose()
